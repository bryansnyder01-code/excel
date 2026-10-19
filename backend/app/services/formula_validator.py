"""Static checks and safe auto-corrections for a WorkbookSpec before it is built.

This is the "refinement & error handling" layer. It never calls the LLM itself;
it returns `Issue`s, and the engine decides whether to send errors back to the
model for a repair pass.

Errors (sent back to the model for repair):
  * formula syntax: unbalanced parentheses or quotes
  * a reference to a sheet that does not exist
  * a cell that references itself (direct circular reference)
  * division by a literal zero

Auto-corrections (applied in place, reported as fixed warnings):
  * invalid / duplicate / over-long sheet names are sanitised, and formulas that
    referenced the old name are rewritten
  * a division whose divisor is a cell, range or expression is wrapped in
    IFERROR(...,0) so an empty or zero input shows 0 instead of #DIV/0!

Warnings (reported, left alone):
  * functions that are not in the known Excel function list
  * references to cells beyond the populated area of the target sheet
  * rows with more cells than the sheet has columns
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl.utils import column_index_from_string, get_column_letter

from app.schemas import Issue, WorkbookSpec

INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME = 31

# Sheet prefix: 'Any name'! or SimpleName!
_SHEET = r"(?:'(?P<qsheet>(?:[^']|'')+)'|(?P<sheet>[A-Za-z_][\w\.]*))!"
_CELL = r"\$?[A-Z]{1,3}\$?[0-9]+"
_COL = r"\$?[A-Z]{1,3}"
REF_RE = re.compile(
    rf"(?<![A-Za-z0-9_\.])(?:{_SHEET})?"
    rf"(?P<ref>{_CELL}(?::{_CELL})?|{_COL}:{_COL})"
    r"(?![A-Za-z0-9_\(])"
)
FUNC_RE = re.compile(r"(?<![A-Za-z0-9_\.])([A-Z][A-Z0-9\.]*)\s*\(")
STRING_RE = re.compile(r'"(?:[^"]|"")*"')

KNOWN_FUNCTIONS = {
    # math & aggregation
    "SUM", "SUMIF", "SUMIFS", "SUMPRODUCT", "AVERAGE", "AVERAGEIF", "AVERAGEIFS",
    "MIN", "MAX", "MINIFS", "MAXIFS", "COUNT", "COUNTA", "COUNTBLANK", "COUNTIF",
    "COUNTIFS", "ROUND", "ROUNDUP", "ROUNDDOWN", "INT", "ABS", "MOD", "POWER",
    "SQRT", "PRODUCT", "MEDIAN", "STDEV", "STDEV.S", "STDEV.P", "VAR", "VAR.S",
    "LARGE", "SMALL", "RANK", "RANK.EQ", "CEILING", "FLOOR", "SUBTOTAL", "LN",
    "LOG", "LOG10", "EXP", "PI", "RAND", "RANDBETWEEN", "PERCENTILE", "QUARTILE",
    # logic
    "IF", "IFS", "IFERROR", "IFNA", "AND", "OR", "NOT", "XOR", "SWITCH", "TRUE", "FALSE",
    # lookup
    "VLOOKUP", "HLOOKUP", "XLOOKUP", "LOOKUP", "INDEX", "MATCH", "XMATCH", "CHOOSE",
    "OFFSET", "INDIRECT", "ROW", "ROWS", "COLUMN", "COLUMNS", "TRANSPOSE",
    # text
    "CONCAT", "CONCATENATE", "TEXTJOIN", "LEFT", "RIGHT", "MID", "LEN", "UPPER",
    "LOWER", "PROPER", "TRIM", "TEXT", "VALUE", "SUBSTITUTE", "REPLACE", "FIND",
    "SEARCH", "REPT", "EXACT",
    # date
    "TODAY", "NOW", "DATE", "YEAR", "MONTH", "DAY", "EDATE", "EOMONTH", "DATEDIF",
    "NETWORKDAYS", "WORKDAY", "WEEKDAY", "WEEKNUM", "DAYS",
    # finance
    "PMT", "IPMT", "PPMT", "FV", "PV", "NPV", "IRR", "XNPV", "XIRR", "RATE", "NPER",
    # info
    "ISBLANK", "ISNUMBER", "ISTEXT", "ISERROR", "ISNA", "NA",
    # SheetCraft's own =AI(prompt, inputs...): evaluated in the app, written as its last value in downloads
    "AI",
}


@dataclass
class _Ref:
    sheet: str | None  # None means "same sheet"
    min_col: int
    min_row: int | None  # None for whole-column refs
    max_col: int
    max_row: int | None


def _strip_strings(formula: str) -> str:
    """Blank out string literals so their contents are not parsed as references."""
    return STRING_RE.sub(lambda m: '"' + " " * (len(m.group(0)) - 2) + '"', formula)


def _parse_cell(token: str) -> tuple[int, int]:
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?([0-9]+)", token)
    assert m, token
    return column_index_from_string(m.group(1)), int(m.group(2))


def iter_refs(formula: str) -> list[_Ref]:
    refs: list[_Ref] = []
    for m in REF_RE.finditer(_strip_strings(formula)):
        sheet = m.group("qsheet")
        sheet = sheet.replace("''", "'") if sheet else m.group("sheet")
        ref = m.group("ref")
        start, _, end = ref.partition(":")
        end = end or start
        if re.fullmatch(_COL, start):  # whole-column, e.g. B:B
            c1 = column_index_from_string(start.replace("$", ""))
            c2 = column_index_from_string(end.replace("$", ""))
            refs.append(_Ref(sheet, min(c1, c2), None, max(c1, c2), None))
        else:
            c1, r1 = _parse_cell(start)
            c2, r2 = _parse_cell(end)
            refs.append(_Ref(sheet, min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2)))
    return refs


def _balanced(formula: str) -> str | None:
    if formula.count('"') % 2:
        return "unbalanced double quotes"
    depth = 0
    for ch in _strip_strings(formula):
        depth += ch == "("
        depth -= ch == ")"
        if depth < 0:
            return "a closing parenthesis has no matching opening one"
    return "unclosed parenthesis" if depth else None


def _divisions(formula: str) -> list[str]:
    """Return the divisor token after each '/' (ignoring string literals)."""
    body = _strip_strings(formula)
    return [m.group(1) for m in re.finditer(r"/\s*([^\s+\-*/^&=<>,;)]+|\()", body)]


def _wrap_iferror(formula: str) -> str:
    return f"=IFERROR({formula[1:]},0)"


def sanitize_sheet_name(name: str, taken: set[str]) -> str:
    clean = INVALID_SHEET_CHARS.sub("-", name).strip().strip("'") or "Sheet"
    clean = clean[:MAX_SHEET_NAME]
    candidate, n = clean, 2
    while candidate.lower() in taken:
        suffix = f" ({n})"
        candidate = clean[: MAX_SHEET_NAME - len(suffix)] + suffix
        n += 1
    return candidate


def _rename_in_formula(formula: str, old: str, new: str) -> str:
    quoted_new = f"'{new.replace(chr(39), chr(39) * 2)}'!"
    formula = formula.replace(f"'{old.replace(chr(39), chr(39) * 2)}'!", quoted_new)
    return re.sub(rf"(?<![A-Za-z0-9_\.']){re.escape(old)}!", quoted_new, formula)


def validate_workbook(spec: WorkbookSpec, autocorrect: bool = True) -> list[Issue]:
    """Validate `spec` in place, applying safe fixes when `autocorrect` is set."""
    issues: list[Issue] = []

    if not spec.sheets:
        return [Issue(severity="error", message="The workbook has no sheets.")]

    # 1. Sheet names
    renames: dict[str, str] = {}
    taken: set[str] = set()
    for sheet in spec.sheets:
        new = sanitize_sheet_name(sheet.name, taken)
        taken.add(new.lower())
        if new != sheet.name:
            issues.append(
                Issue(severity="warning", sheet=new, fixed=autocorrect,
                      message=f'Sheet name "{sheet.name}" is not valid in Excel; renamed to "{new}".')
            )
            if autocorrect:
                renames[sheet.name] = new
                sheet.name = new
    if renames:
        for sheet in spec.sheets:
            for row in sheet.rows:
                for cell in row.cells:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        for old, new in renames.items():
                            cell.value = _rename_in_formula(cell.value, old, new)
        _rename_extras(spec, renames)

    # Populated extent of each sheet: (columns, last row incl. header)
    extent = {s.name.lower(): (len(s.columns), len(s.rows) + 1) for s in spec.sheets}

    # 2. Cells
    for sheet in spec.sheets:
        ncols = len(sheet.columns)
        for r_idx, row in enumerate(sheet.rows):
            excel_row = r_idx + 2
            if len(row.cells) > ncols:
                issues.append(
                    Issue(severity="warning", sheet=sheet.name, cell=f"row {excel_row}",
                          message=f"Row has {len(row.cells)} cells but the sheet has {ncols} columns.")
                )
            for c_idx, cell in enumerate(row.cells):
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                addr = f"{get_column_letter(c_idx + 1)}{excel_row}"
                issues.extend(
                    _check_formula(cell, sheet.name, addr, c_idx + 1, excel_row, extent, autocorrect)
                )

    # 3. Chart and data-validation ranges
    for sheet in spec.sheets:
        for chart in sheet.charts:
            ranges = [("categories", chart.categories)] + [(f'series "{s.name}"', s.range) for s in chart.series]
            if not chart.series:
                issues.append(Issue(severity="error", sheet=sheet.name, message=f'Chart "{chart.title}" has no data series.'))
            for label, rng in ranges:
                if not _valid_range(rng):
                    issues.append(Issue(severity="error", sheet=sheet.name,
                                        message=f'Chart "{chart.title}" {label} range "{rng}" is not a valid A1 range.'))
            _check_chart(chart, sheet.name, issues)
        for dv in sheet.validations:
            if not _valid_range(dv.range):
                issues.append(Issue(severity="error", sheet=sheet.name,
                                    message=f'Validation range "{dv.range}" is not a valid A1 range.'))
            elif dv.type.value == "list" and not [o for o in dv.options if o.strip()]:
                issues.append(Issue(severity="warning", sheet=sheet.name, cell=dv.range,
                                    message="A dropdown has no options, so it was skipped."))
        _check_layout(sheet, issues)
        _check_protection(sheet, issues)
    _check_dashboard(spec, extent, issues)
    return issues


def _rename_extras(spec: WorkbookSpec, renames: dict[str, str]) -> None:
    """Dashboard references and connection targets follow a renamed sheet."""
    for old, new in renames.items():
        if spec.dashboard:
            for tile in spec.dashboard.tiles:
                for attr in ("ref", "compare_ref"):
                    v = getattr(tile, attr)
                    if v:
                        setattr(tile, attr, _rename_in_formula(v, old, new))
                if tile.sheet == old:
                    tile.sheet = new
        for conn in spec.connections:
            if conn.sheet == old:
                conn.sheet = new


def _check_protection(sheet, issues: list[Issue]) -> None:
    for rng in sheet.locked_ranges:
        if not _valid_range(rng):
            issues.append(Issue(severity="warning", sheet=sheet.name, cell=rng,
                                message=f'Locked range "{rng}" is not a valid A1 range like "A1:H1", so it was skipped.'))


def _split_ref(ref: str) -> tuple[str | None, str]:
    sheet, sep, cell = ref.strip().lstrip("=").rpartition("!")
    if not sep:
        return None, cell.strip()
    sheet = sheet.strip()
    if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, cell.strip()


def _check_dashboard(spec: WorkbookSpec, extent: dict, issues: list[Issue]) -> None:
    """Dashboard tiles that point at sheets, cells or charts that don't exist: reported, the tile shows #REF!."""
    if not spec.dashboard:
        return
    sheets = {s.name.lower(): s for s in spec.sheets}
    for tile in spec.dashboard.tiles:
        label = f'Dashboard tile "{tile.title or tile.id}"'
        if tile.kind.value == "kpi":
            for attr in ("ref", "compare_ref"):
                ref = getattr(tile, attr)
                if not ref:
                    if attr == "ref":
                        issues.append(Issue(severity="warning", message=f"{label} has no cell to show."))
                    continue
                sheet, cell = _split_ref(ref)
                if sheet is None or sheet.lower() not in sheets:
                    issues.append(Issue(severity="warning", sheet=sheet, cell=cell,
                                        message=f'{label} points at "{ref}", but that sheet does not exist.'))
                elif not re.fullmatch(_CELL, cell.upper()):
                    issues.append(Issue(severity="warning", sheet=sheet, cell=cell,
                                        message=f'{label} points at "{ref}", which is not a single cell like "B4".'))
                else:
                    col, row = _parse_cell(cell.upper())
                    ncols, nrows = extent[sheet.lower()]
                    if col > ncols or row > nrows:
                        issues.append(Issue(severity="warning", sheet=sheet, cell=cell.upper(),
                                            message=f'{label} points at "{ref}", which is outside the sheet\'s data.'))
        elif tile.kind.value == "chart":
            sheet = sheets.get((tile.sheet or "").lower())
            if sheet is None:
                issues.append(Issue(severity="warning", sheet=tile.sheet,
                                    message=f'{label} shows a chart from sheet "{tile.sheet}", which does not exist.'))
            elif not 0 <= (tile.chart or 0) < len(sheet.charts):
                issues.append(Issue(severity="warning", sheet=sheet.name,
                                    message=f'{label} shows chart {(tile.chart or 0) + 1} of "{sheet.name}", which has {len(sheet.charts)}.'))


_HEX_RE = re.compile(r"^#?([0-9A-Fa-f]{3}|[0-9A-Fa-f]{6})$")


def _check_chart(chart, sheet_name: str, issues: list[Issue]) -> None:
    kind = chart.type.value
    if kind in ("pie", "doughnut") and len(chart.series) > 1:
        issues.append(Issue(severity="warning", sheet=sheet_name,
                            message=f'Chart "{chart.title}" is a {kind} chart, so only its first series is shown.'))
    if chart.stacked and kind not in ("column", "bar", "area"):
        issues.append(Issue(severity="warning", sheet=sheet_name,
                            message=f'Chart "{chart.title}": only column, bar and area charts can be stacked.'))
    for s in chart.series:
        if s.kind is not None and s.kind.value == "line" and kind != "column":
            issues.append(Issue(severity="warning", sheet=sheet_name,
                                message=f'Chart "{chart.title}": series "{s.name}" can only be drawn as a line on a column chart.'))
        if s.color and not _HEX_RE.match(s.color.strip()):
            issues.append(Issue(severity="warning", sheet=sheet_name,
                                message=f'Chart "{chart.title}": series colour "{s.color}" is not a hex colour like "#4472C4".'))


def _check_layout(sheet, issues: list[Issue]) -> None:
    """Merged ranges and column filters: reported, never fatal (the builder skips bad ones)."""
    boxes: list[tuple[int, int, int, int]] = []
    for rng in sheet.merges:
        if not _valid_range(rng) or ":" not in rng:
            issues.append(Issue(severity="warning", sheet=sheet.name, cell=rng,
                                message=f'Merged range "{rng}" is not a valid A1 range like "A1:C1", so it was skipped.'))
            continue
        a, b = (_parse_cell(p) for p in rng.strip().upper().split(":"))
        box = (min(a[0], b[0]), min(a[1], b[1]), max(a[0], b[0]), max(a[1], b[1]))
        if any(box[0] <= t[2] and t[0] <= box[2] and box[1] <= t[3] and t[1] <= box[3] for t in boxes):
            issues.append(Issue(severity="warning", sheet=sheet.name, cell=rng,
                                message=f'Merged range "{rng}" overlaps another merge, so it was skipped.'))
            continue
        boxes.append(box)
    ncols = len(sheet.columns)
    for f in sheet.filters:
        if not 0 <= f.column < ncols:
            issues.append(Issue(severity="warning", sheet=sheet.name,
                                message=f"A filter points at column {f.column + 1}, but the sheet has {ncols} columns, so it was skipped."))


_RANGE_RE = re.compile(r"^\$?[A-Z]{1,3}\$?[0-9]+(:\$?[A-Z]{1,3}\$?[0-9]+)?$")


def _valid_range(a1: str) -> bool:
    return bool(_RANGE_RE.match(a1.strip().upper()))


def _check_formula(cell, sheet_name, addr, col, row, extent, autocorrect) -> list[Issue]:
    formula: str = cell.value
    out: list[Issue] = []

    def issue(sev: str, msg: str, fixed: bool = False) -> None:
        out.append(Issue(severity=sev, sheet=sheet_name, cell=addr, message=msg, fixed=fixed))

    problem = _balanced(formula)
    if problem:
        issue("error", f"Formula {formula} has {problem}.")
        return out

    for fn in FUNC_RE.findall(_strip_strings(formula)):
        if fn not in KNOWN_FUNCTIONS and not fn.startswith("_XLFN."):
            issue("warning", f"{fn} is not a recognised Excel function.")

    for ref in iter_refs(formula):
        target = (ref.sheet or sheet_name).lower()
        if target not in extent:
            issue("error", f'Formula {formula} references sheet "{ref.sheet}", which does not exist.')
            continue
        same_sheet = target == sheet_name.lower()
        if (
            same_sheet
            and ref.min_col <= col <= ref.max_col
            and (ref.min_row is None or ref.min_row <= row <= ref.max_row)
        ):
            issue("error", f"Formula {formula} refers to its own cell (circular reference).")
        ncols, nrows = extent[target]
        if ref.max_col > ncols or (ref.max_row is not None and ref.max_row > nrows):
            issue("warning", f"Formula {formula} reaches past the populated area of the sheet.")

    divisors = _divisions(formula)
    if any(re.fullmatch(r"0+(\.0+)?", d) for d in divisors):
        issue("error", f"Formula {formula} divides by zero.")
    elif autocorrect and divisors and not formula.upper().startswith(("=IFERROR(", "=IF(")):
        needs_guard = any(not re.fullmatch(r"[0-9.]+%?", d) for d in divisors)
        if needs_guard:
            cell.value = _wrap_iferror(formula)
            issue("warning", "Wrapped a division in IFERROR so an empty or zero divisor shows 0.", fixed=True)
    return out


def has_errors(issues: list[Issue]) -> bool:
    return any(i.severity == "error" and not i.fixed for i in issues)
