"""Compile a WorkbookSpec into a styled .xlsx with live Excel formulas (openpyxl).

Formulas are written as formulas, never as computed values, so the downloaded
file recalculates in Excel, LibreOffice and Google Sheets.
"""

from __future__ import annotations

import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import column_index_from_string, get_column_letter, quote_sheetname, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.schemas import (
    CURRENCY_SYMBOLS,
    BorderStyle,
    CellSpec,
    ChartSpec,
    ColumnSpec,
    ChartType,
    DashboardSpec,
    DashboardTile,
    DataValidationSpec,
    ValidationType,
    CondOperator,
    CondStyle,
    NumberFormat,
    RowKind,
    SeriesKind,
    SheetSpec,
    WorkbookSpec,
)

# ---------------------------------------------------------------------------
# Corporate theme
# ---------------------------------------------------------------------------
FONT_NAME = "Calibri"
ACCENT = "1F4E78"  # dark blue header fill
ACCENT_LIGHT = "DDEBF7"  # subtotal / section fill
TOTAL_FILL = "BDD7EE"
BAND_FILL = "F7F9FC"  # zebra banding on data rows
BORDER_COLOR = "BFBFBF"

NUMBER_FORMATS = {
    NumberFormat.text: "@",
    NumberFormat.integer: "#,##0",
    NumberFormat.decimal: "#,##0.00",
    NumberFormat.currency: '"$"#,##0.00;[Red]-"$"#,##0.00',
    NumberFormat.percent: "0.0%",
    NumberFormat.date: "yyyy-mm-dd",
}

# Default decimal places when a column / cell doesn't set `decimals`.
DEFAULT_DECIMALS = {NumberFormat.decimal: 2, NumberFormat.currency: 2, NumberFormat.percent: 1}
ZERO_DECIMAL_CURRENCIES = {"JPY"}

COND_FILLS = {
    CondStyle.good: ("C6EFCE", "006100"),
    CondStyle.bad: ("FFC7CE", "9C0006"),
    CondStyle.warning: ("FFEB9C", "9C5700"),
}

# Colour scale (low -> mid -> high) and the default data bar colour.
SCALE_LOW, SCALE_MID, SCALE_HIGH = "F8696B", "FFEB84", "63BE7B"
DATA_BAR = "638EC6"

THIN = Side(style="thin", color=BORDER_COLOR)
MEDIUM = Side(style="medium", color=ACCENT)
GRID_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(left=THIN, right=THIN, top=MEDIUM, bottom=Side(style="double", color=ACCENT))

HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
BODY_FONT = Font(name=FONT_NAME, size=11)
BOLD_FONT = Font(name=FONT_NAME, size=11, bold=True)

MIN_WIDTH, MAX_WIDTH = 8, 60
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_HEX = re.compile(r"^[0-9A-Fa-f]{6}$")


def number_format(fmt: NumberFormat, decimals: int | None = None, currency: str | None = None) -> str:
    """Excel number format for a column / cell format, honouring decimal places and currency code."""
    if fmt not in DEFAULT_DECIMALS:
        return NUMBER_FORMATS[fmt]
    code = (currency or "USD").strip().upper()
    if code not in CURRENCY_SYMBOLS:
        code = "USD"
    if decimals is None:
        decimals = 0 if fmt == NumberFormat.currency and code in ZERO_DECIMAL_CURRENCIES else DEFAULT_DECIMALS[fmt]
    decimals = max(0, min(10, int(decimals)))
    frac = "." + "0" * decimals if decimals else ""
    if fmt == NumberFormat.percent:
        return f"0{frac}%"
    if fmt == NumberFormat.decimal:
        return f"#,##0{frac}"
    sym = '"' + CURRENCY_SYMBOLS[code].replace('"', "") + '"'
    return f"{sym}#,##0{frac};[Red]-{sym}#,##0{frac}"


def _cell_number_format(spec: CellSpec, col: ColumnSpec | None) -> tuple[NumberFormat, str]:
    """The effective format for a cell (cell settings override the column's) and its Excel format string."""
    fmt = spec.format or (col.format if col else NumberFormat.text)
    decimals = spec.decimals if spec.decimals is not None else (col.decimals if col else None)
    currency = spec.currency or (col.currency if col else None)
    return fmt, number_format(fmt, decimals, currency)


def build_workbook(spec: WorkbookSpec) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = spec.title
    wb.properties.creator = "SheetCraft AI"
    for sheet in spec.sheets:
        _build_sheet(wb.create_sheet(title=sheet.name), sheet)
    for sheet in spec.sheets:
        ws = wb[sheet.name]
        for dv in sheet.validations:
            _add_validation(wb, ws, dv)
        for i, chart in enumerate(sheet.charts):
            _add_chart(ws, sheet, chart, i)
        _apply_protection(ws, sheet)
    if spec.scenarios:
        _add_scenarios_sheet(wb, spec)
    if spec.dashboard and spec.dashboard.tiles:
        _add_dashboard_sheet(wb, spec, spec.dashboard)
    return wb


def workbook_to_bytes(spec: WorkbookSpec) -> bytes:
    buf = io.BytesIO()
    build_workbook(spec).save(buf)
    return buf.getvalue()


def _build_sheet(ws: Worksheet, sheet: SheetSpec) -> None:
    ncols = len(sheet.columns)

    # Header row
    for c, col in enumerate(sheet.columns, start=1):
        cell = ws.cell(row=1, column=c, value=col.header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = GRID_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = sheet.header_height or 22

    # Body rows
    for r_idx, row in enumerate(sheet.rows):
        excel_row = r_idx + 2
        banded = row.kind == RowKind.data and r_idx % 2 == 1
        width = max(ncols, len(row.cells))
        for c in range(1, width + 1):
            spec_cell = row.cells[c - 1] if c - 1 < len(row.cells) else CellSpec()
            col = sheet.columns[c - 1] if c - 1 < ncols else None
            _write_cell(ws, excel_row, c, spec_cell, col, row.kind, banded)
        if row.height:
            ws.row_dimensions[excel_row].height = row.height

    ws.freeze_panes = freeze_cell(sheet)
    ws.sheet_view.showGridLines = sheet.show_gridlines
    if ncols and sheet.rows:
        last_data = _last_data_row(sheet)
        if last_data >= 2:
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last_data}"
            _apply_filters(ws, sheet)
    _apply_merges(ws, sheet)

    for cf in sheet.conditional_formats:
        _add_conditional_format(ws, cf)

    _fit_columns(ws, sheet)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if ncols > 6 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def freeze_cell(sheet: SheetSpec) -> str | None:
    """Top-left unfrozen cell: header row and/or leading columns stay visible ("C2" = row 1 + columns A:B)."""
    cols = max(0, min(sheet.freeze_columns or 0, 50))
    row = 2 if sheet.freeze_header else 1
    if cols == 0 and row == 1:
        return None
    return f"{get_column_letter(cols + 1)}{row}"


_AI_CALL = re.compile(r"(?<![A-Za-z0-9_.])AI\s*\(", re.I)
_STRING = re.compile(r'"(?:[^"]|"")*"')


def is_ai_formula(value) -> bool:
    """True for a formula that calls SheetCraft's =AI(prompt, inputs...) function (outside string literals)."""
    return isinstance(value, str) and value.startswith("=") and bool(_AI_CALL.search(_STRING.sub('""', value)))


def _write_cell(ws, row, col, spec: CellSpec, col_spec: ColumnSpec | None, kind: RowKind, banded: bool) -> None:
    value = spec.value
    if is_ai_formula(value):
        # Excel has no AI function: write the last result as a plain value (empty if never computed).
        value = spec.ai_value if spec.ai_value not in (None, "") else None
        if isinstance(value, str) and value.startswith("="):
            value = "'" + value  # never let an AI answer become a formula
    fmt, excel_fmt = _cell_number_format(spec, col_spec)
    if isinstance(value, str) and fmt == NumberFormat.date and _ISO_DATE.match(value):
        value = date.fromisoformat(value)
    elif isinstance(value, str) and not value.startswith("=") and fmt != NumberFormat.text:
        value = _coerce_number(value)

    cell = ws.cell(row=row, column=col, value=value)
    is_formula = isinstance(value, str) and value.startswith("=")
    is_number = isinstance(value, (int, float, date, datetime)) and not isinstance(value, bool)

    if is_formula or is_number:
        # A formula or number in a "text" column keeps Excel's General format so it still calculates.
        cell.number_format = "General" if fmt == NumberFormat.text else excel_fmt
        cell.alignment = Alignment(horizontal="right", vertical="center")
    else:
        cell.number_format = "@" if fmt == NumberFormat.text else excel_fmt
        cell.alignment = Alignment(horizontal="left", vertical="center")

    bold = spec.bold or kind in (RowKind.total, RowKind.subtotal, RowKind.section)
    cell.font = BOLD_FONT if bold else BODY_FONT
    cell.border = TOTAL_BORDER if kind == RowKind.total else GRID_BORDER
    if kind == RowKind.total:
        cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)
    elif kind in (RowKind.subtotal, RowKind.section):
        cell.fill = PatternFill("solid", fgColor=ACCENT_LIGHT)
    elif banded:
        cell.fill = PatternFill("solid", fgColor=BAND_FILL)

    _apply_overrides(cell, spec, bold)

    if spec.note:
        cell.comment = Comment(spec.note, "SheetCraft AI")


def _hex(color: str | None) -> str | None:
    """'#c00000' or 'C00000' -> 'C00000'; anything else -> None."""
    if not color:
        return None
    c = color.strip().lstrip("#")
    if len(c) == 3:
        c = "".join(ch * 2 for ch in c)
    return c.upper() if _HEX.match(c) else None


def _apply_overrides(cell, spec: CellSpec, bold: bool) -> None:
    """Layer the user's / model's explicit styling over the automatic theme."""
    font_color = _hex(spec.font_color)
    if spec.italic or spec.underline or font_color:
        cell.font = Font(
            name=FONT_NAME, size=11, bold=bold, italic=spec.italic,
            underline="single" if spec.underline else None,
            color=font_color or cell.font.color,
        )
    fill = _hex(spec.fill_color)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if spec.align:
        cell.alignment = Alignment(horizontal=str(getattr(spec.align, "value", spec.align)), vertical="center")
    if spec.wrap:
        a = cell.alignment
        cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical or "center", wrap_text=True)
    if spec.border == BorderStyle.none:
        cell.border = Border()
    elif spec.border == BorderStyle.thick:
        side = Side(style="medium", color="000000")
        cell.border = Border(left=side, right=side, top=side, bottom=side)
    elif spec.border == BorderStyle.thin:
        cell.border = GRID_BORDER


def _coerce_number(value: str):
    """Turn '$1,200', '15%' or '3.5' into numbers; leave anything else as text."""
    s = value.strip().replace(",", "")
    for sym in sorted(set(CURRENCY_SYMBOLS.values()), key=len, reverse=True):
        s = s.replace(sym.strip(), "")
    s = s.strip()
    try:
        return float(s[:-1]) / 100 if s.endswith("%") else float(s)
    except ValueError:
        return value


def _last_data_row(sheet: SheetSpec) -> int:
    last = 1
    for i, row in enumerate(sheet.rows):
        if row.kind == RowKind.data:
            last = i + 2
    return last


def _add_conditional_format(ws: Worksheet, cf) -> None:
    if cf.operator == CondOperator.color_scale:
        low, high = (SCALE_HIGH, SCALE_LOW) if cf.reverse else (SCALE_LOW, SCALE_HIGH)
        rule = ColorScaleRule(start_type="min", start_color=low, mid_type="percentile", mid_value=50,
                              mid_color=SCALE_MID, end_type="max", end_color=high)
        ws.conditional_formatting.add(cf.range, rule)
        return
    if cf.operator == CondOperator.data_bar:
        rule = DataBarRule(start_type="min", end_type="max", color=_hex(cf.fill_color) or DATA_BAR)
        ws.conditional_formatting.add(cf.range, rule)
        return

    if cf.style == CondStyle.custom:
        fill_hex, font_hex = _hex(cf.fill_color), _hex(cf.font_color)
        if not fill_hex and not font_hex:
            fill_hex, font_hex = COND_FILLS[CondStyle.good]
    else:
        fill_hex, font_hex = COND_FILLS[cf.style]
    style = dict(
        fill=PatternFill("solid", fgColor=fill_hex, bgColor=fill_hex) if fill_hex else None,
        font=Font(color=font_hex) if font_hex else None,
    )
    if cf.operator == CondOperator.contains:
        if not cf.text:
            return
        min_col, min_row, *_ = range_boundaries(cf.range.split()[0].replace("$", ""))
        text = cf.text.replace('"', '""')
        rule = FormulaRule(formula=[f'ISNUMBER(SEARCH("{text}",{get_column_letter(min_col)}{min_row}))'], **style)
    elif cf.operator == CondOperator.negative:
        rule = CellIsRule(operator="lessThan", formula=["0"], **style)
    elif cf.operator == CondOperator.between:
        lo, hi = cf.value1 or 0, cf.value2 if cf.value2 is not None else cf.value1 or 0
        rule = CellIsRule(operator="between", formula=[str(lo), str(hi)], **style)
    else:
        op = {"greater_than": "greaterThan", "less_than": "lessThan", "equal": "equal"}[cf.operator.value]
        rule = CellIsRule(operator=op, formula=[str(cf.value1 or 0)], **style)
    ws.conditional_formatting.add(cf.range, rule)


def _display_len(value, fmt: NumberFormat, excel_fmt: str = "") -> int:
    if value is None:
        return 0
    if isinstance(value, str) and value.startswith("="):
        return 14 if fmt == NumberFormat.currency else 12  # formulas: assume a typical result width
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        positive = excel_fmt.split(";")[0]
        decimals = len(positive.split(".")[1].rstrip("%")) if "." in positive else 0
        if fmt == NumberFormat.currency:
            symbol = positive.split('"')[1] if positive.count('"') >= 2 else "$"
            return len(f"{symbol}{value:,.{decimals}f}") + 1
        if fmt == NumberFormat.percent:
            return len(f"{value * 100:.{decimals}f}%")
        return len(f"{value:,.2f}")
    return max(len(line) for line in str(value).splitlines() or [""])


def _fit_columns(ws: Worksheet, sheet: SheetSpec) -> None:
    for c, col in enumerate(sheet.columns, start=1):
        letter = get_column_letter(c)
        if col.width:
            ws.column_dimensions[letter].width = col.width
            continue
        longest = len(col.header) + 2  # room for the filter arrow
        for row in sheet.rows:
            if c - 1 < len(row.cells):
                cell = row.cells[c - 1]
                fmt, excel_fmt = _cell_number_format(cell, col)
                longest = max(longest, _display_len(cell.value, fmt, excel_fmt))
        ws.column_dimensions[letter].width = max(MIN_WIDTH, min(MAX_WIDTH, longest + 2))


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------
CHART_W_CM, CHART_H_CM = 16, 8
CHART_ROWS = 17  # rows a chart occupies when stacked


def _ref(ws: Worksheet, a1: str) -> Reference:
    min_col, min_row, max_col, max_row = range_boundaries(a1.replace("$", ""))
    return Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)


def _add_chart(ws: Worksheet, sheet: SheetSpec, spec: ChartSpec, index: int,
               target: Worksheet | None = None, size: tuple[float, float] | None = None) -> None:
    """Draw `spec` from the data on `ws`; `target` places it on another sheet (the dashboard)."""
    single = spec.type in (ChartType.pie, ChartType.doughnut)
    series = spec.series[:1] if single else spec.series
    combo_lines = [s for s in series if s.kind == SeriesKind.line] if spec.type == ChartType.column else []
    main_series = [s for s in series if s not in combo_lines]
    if spec.type == ChartType.column and not main_series:  # every series asked for a line
        main_series, combo_lines = combo_lines, []
        chart = LineChart()
    elif spec.type == ChartType.pie:
        chart = PieChart()
    elif spec.type == ChartType.doughnut:
        chart = DoughnutChart()
    elif spec.type == ChartType.line:
        chart = LineChart()
    elif spec.type == ChartType.scatter:
        chart = ScatterChart()
        chart.scatterStyle = "lineMarker"
    elif spec.type == ChartType.area:
        chart = AreaChart()
    else:
        chart = BarChart()
        chart.type = "bar" if spec.type == ChartType.bar else "col"
        chart.gapWidth = 80
    if spec.stacked and isinstance(chart, (BarChart, AreaChart)):
        chart.grouping = "stacked"
        if isinstance(chart, BarChart):
            chart.overlap = 100
    chart.title = spec.title
    chart.width, chart.height = size or (CHART_W_CM, CHART_H_CM)
    chart.style = 10

    if isinstance(chart, ScatterChart):
        xvalues = _ref(ws, spec.categories)
        for s in main_series:
            ser = Series(_ref(ws, s.range), xvalues=xvalues, title=s.name)
            ser.marker.symbol = "circle"
            ser.marker.size = 7
            ser.graphicalProperties.line.noFill = True  # points only
            color = _hex(s.color)
            if color:
                ser.marker.graphicalProperties.solidFill = color
                ser.marker.graphicalProperties.line.solidFill = color
            chart.series.append(ser)
    else:
        for s in main_series:
            ser = Series(_ref(ws, s.range), title=s.name)
            if not single:
                _color_series(ser, s.color, line=isinstance(chart, LineChart))
            chart.series.append(ser)
        chart.set_categories(_ref(ws, spec.categories))

    if combo_lines:
        line = LineChart()
        for s in combo_lines:
            ser = Series(_ref(ws, s.range), title=s.name)
            _color_series(ser, s.color, line=True)
            line.series.append(ser)
        line.set_categories(_ref(ws, spec.categories))
        chart += line  # same axes: openpyxl draws both chart types in one plot area

    if not single:
        chart.legend.position = "b"
        if len(series) == 1:
            chart.legend = None
        fmt = _series_format(sheet, spec.series[0].range) if spec.series else None
        if fmt:
            chart.y_axis.numFmt = fmt
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        if spec.x_title:
            chart.x_axis.title = spec.x_title
        if spec.y_title:
            chart.y_axis.title = spec.y_title

    anchor = spec.anchor or f"{get_column_letter(len(sheet.columns) + 2)}{2 + index * CHART_ROWS}"
    (target or ws).add_chart(chart, anchor)


def _color_series(ser: Series, color: str | None, line: bool) -> None:
    color = _hex(color)
    if not color:
        return
    if line:
        ser.graphicalProperties.line.solidFill = color
    else:
        ser.graphicalProperties.solidFill = color
        ser.graphicalProperties.line.solidFill = color


def _series_format(sheet: SheetSpec, a1: str) -> str | None:
    min_col, *_ = range_boundaries(a1.replace("$", ""))
    if min_col - 1 < len(sheet.columns):
        col = sheet.columns[min_col - 1]
        if col.format in (NumberFormat.currency, NumberFormat.percent, NumberFormat.integer, NumberFormat.decimal):
            return number_format(col.format, col.decimals, col.currency).split(";")[0]
    return None


# ---------------------------------------------------------------------------
# Data validation (dropdowns, number / date limits)
# ---------------------------------------------------------------------------
LISTS_SHEET = "_lists"


def _add_validation(wb: Workbook, ws: Worksheet, spec: DataValidationSpec) -> None:
    error = spec.message or "Please enter a valid value."
    if spec.type == ValidationType.list:
        options = [o.replace('"', "'") for o in spec.options if o.strip()]
        if not options:
            return
        inline = ",".join(options)
        if len(inline) <= 250 and not any("," in o for o in options):
            formula = f'"{inline}"'
        else:
            # Long or comma-containing lists live on a hidden helper sheet.
            lists = wb[LISTS_SHEET] if LISTS_SHEET in wb.sheetnames else wb.create_sheet(LISTS_SHEET)
            lists.sheet_state = "hidden"
            col = lists.max_column + 1 if lists["A1"].value is not None else 1
            for r, opt in enumerate(options, start=1):
                lists.cell(row=r, column=col, value=opt)
            letter = get_column_letter(col)
            formula = f"{quote_sheetname(LISTS_SHEET)}!${letter}$1:${letter}${len(options)}"
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    else:
        kind = {"whole": "whole", "decimal": "decimal", "date": "date"}[spec.type.value]
        if spec.min is not None and spec.max is not None:
            dv = DataValidation(type=kind, operator="between", formula1=str(spec.min), formula2=str(spec.max))
        elif spec.min is not None:
            dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1=str(spec.min))
        elif spec.max is not None:
            dv = DataValidation(type=kind, operator="lessThanOrEqual", formula1=str(spec.max))
        else:
            dv = DataValidation(type=kind, operator="greaterThanOrEqual", formula1="0" if kind != "date" else "1")
    dv.error, dv.errorTitle, dv.showErrorMessage = error, "Invalid entry", True
    dv.add(spec.range)
    ws.add_data_validation(dv)


# ---------------------------------------------------------------------------
# Merges, filters and scenarios
# ---------------------------------------------------------------------------
SCENARIOS_SHEET = "Scenarios"


def _apply_merges(ws: Worksheet, sheet: SheetSpec) -> None:
    """Merge each valid, non-overlapping range; the top-left cell keeps its value and style."""
    taken: list[tuple[int, int, int, int]] = []
    for rng in sheet.merges:
        try:
            box = range_boundaries(rng.strip().upper().replace("$", ""))
        except (ValueError, TypeError):
            continue
        if None in box or (box[0] == box[2] and box[1] == box[3]):
            continue
        if any(box[0] <= t[2] and t[0] <= box[2] and box[1] <= t[3] and t[1] <= box[3] for t in taken):
            continue
        taken.append(box)
        ws.merge_cells(start_column=box[0], start_row=box[1], end_column=box[2], end_row=box[3])


def display_text(value) -> str:
    """What a literal cell shows, without number formatting: 100.0 -> "100", True -> "TRUE"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value)


def _apply_filters(ws: Worksheet, sheet: SheetSpec) -> None:
    """Write filter criteria and hide the data rows they exclude (formula cells can't be evaluated: left visible)."""
    ncols = len(sheet.columns)
    filters = [f for f in sheet.filters if 0 <= f.column < ncols]
    for f in filters:
        ws.auto_filter.add_filter_column(f.column, list(f.values))
    if not filters:
        return
    for r_idx, row in enumerate(sheet.rows):
        if row.kind != RowKind.data:
            continue
        for f in filters:
            value = row.cells[f.column].value if f.column < len(row.cells) else None
            if isinstance(value, str) and value.startswith("="):
                continue
            allowed = {v.strip().lower() for v in f.values}
            if display_text(value).strip().lower() not in allowed:
                ws.row_dimensions[r_idx + 2].hidden = True
                break


def _scenario_address(address: str) -> str:
    """"Inputs!$b$2" or "'Inputs'!B2" -> "'Inputs'!B2"."""
    sheet, sep, cell = address.strip().rpartition("!")
    cell = cell.replace("$", "").upper()
    if not sep:
        return cell
    sheet = sheet.strip()
    if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return "'" + sheet.replace("'", "''") + "'!" + cell


def _add_scenarios_sheet(wb: Workbook, spec: WorkbookSpec) -> None:
    """A plain table: one row per input cell, one column per scenario."""
    addresses: list[str] = []
    for sc in spec.scenarios:
        for addr in sc.values:
            norm = _scenario_address(addr)
            if norm not in addresses:
                addresses.append(norm)
    rows = []
    for addr in addresses:
        cells = [CellSpec(value=addr)]
        for sc in spec.scenarios:
            value = next((v for a, v in sc.values.items() if _scenario_address(a) == addr), None)
            cells.append(CellSpec(value=value))
        rows.append({"cells": cells})
    name, n = SCENARIOS_SHEET, 2
    while name.lower() in {s.lower() for s in wb.sheetnames}:
        name, n = f"{SCENARIOS_SHEET} ({n})", n + 1
    sheet = SheetSpec(
        name=name,
        columns=[ColumnSpec(header="Input cell")] + [ColumnSpec(header=sc.name) for sc in spec.scenarios],
        rows=rows,
    )
    _build_sheet(wb.create_sheet(title=name), sheet)
    wb[name].auto_filter.ref = None


# ---------------------------------------------------------------------------
# Sheet protection
# ---------------------------------------------------------------------------
UNLOCKED = Protection(locked=False)
LOCKED = Protection(locked=True)


def _protect(ws: Worksheet) -> None:
    ws.protection.sheet = True
    ws.protection.autoFilter = False  # filtering and sorting still work on a protected sheet
    ws.protection.sort = False
    ws.protection.formatColumns = False


def _apply_protection(ws: Worksheet, sheet: SheetSpec) -> None:
    """locked: the whole sheet is read-only. locked_ranges: only those cells are; everything else is unlocked."""
    if sheet.locked:
        _protect(ws)
        return
    boxes = []
    for rng in sheet.locked_ranges:
        try:
            box = range_boundaries(rng.strip().upper().replace("$", ""))
        except (ValueError, TypeError):
            continue
        if None in box:
            continue
        boxes.append(box)
    if not boxes:
        return
    _protect(ws)
    max_col = max([ws.max_column] + [b[2] for b in boxes])
    max_row = max([ws.max_row] + [b[3] for b in boxes])
    for c in range(1, max_col + 1):  # new cells typed into these columns stay editable too
        ws.column_dimensions[get_column_letter(c)].protection = UNLOCKED
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(row=r, column=c).protection = UNLOCKED
    for min_col, min_row, bmax_col, bmax_row in boxes:
        for r in range(min_row, bmax_row + 1):
            for c in range(min_col, bmax_col + 1):
                ws.cell(row=r, column=c).protection = LOCKED


# ---------------------------------------------------------------------------
# Dashboard sheet
# ---------------------------------------------------------------------------
DASHBOARD_SHEET = "Dashboard"
DASH_COL_WIDTH = 14  # characters per grid unit (12 units across)
DASH_ROWS_PER_UNIT = 3
DASH_FIRST_ROW = 3  # row 1: title, row 2: spacer
_CM_PER_CHAR, _CM_PER_ROW = 0.19, 0.53
KPI_LABEL_FONT = Font(name=FONT_NAME, size=10, color="595959")
KPI_VALUE_FONT = Font(name=FONT_NAME, size=20, bold=True, color=ACCENT)
KPI_NOTE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")
TILE_FILL = PatternFill("solid", fgColor="F2F6FB")


def parse_cell_ref(ref: str | None, spec: WorkbookSpec) -> tuple[SheetSpec, int, int] | None:
    """'Summary!B4' / "'My Sheet'!$B$4" -> (sheet, row, col), or None if the sheet or address is unknown."""
    if not ref or not isinstance(ref, str):
        return None
    sheet_part, sep, cell = ref.strip().lstrip("=").rpartition("!")
    if not sep:
        return None
    sheet_part = sheet_part.strip()
    if len(sheet_part) >= 2 and sheet_part[0] == sheet_part[-1] == "'":
        sheet_part = sheet_part[1:-1].replace("''", "'")
    sheet = next((s for s in spec.sheets if s.name.lower() == sheet_part.lower()), None)
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([0-9]+)", cell.strip())
    if sheet is None or not m:
        return None
    col, row = column_index_from_string(m.group(1).upper()), int(m.group(2))
    if row < 1 or col > 16384 or row > 1048576:
        return None
    return sheet, row, col


def _source_format(sheet: SheetSpec, row: int, col: int) -> str:
    """The Excel number format of a cell in the spec ("General" for headers, text and unknown cells)."""
    if row < 2:
        return "General"
    col_spec = sheet.columns[col - 1] if col - 1 < len(sheet.columns) else None
    cells = sheet.rows[row - 2].cells if row - 2 < len(sheet.rows) else []
    cell = cells[col - 1] if col - 1 < len(cells) else CellSpec()
    fmt, excel_fmt = _cell_number_format(cell, col_spec)
    return "General" if fmt == NumberFormat.text else excel_fmt


def _ref_formula(sheet: SheetSpec, row: int, col: int) -> str:
    """="'Summary'!B4"."""
    return "='" + sheet.name.replace("'", "''") + f"'!{get_column_letter(col)}{row}"


def _add_dashboard_sheet(wb: Workbook, spec: WorkbookSpec, dash: DashboardSpec) -> None:
    """A first tab that mirrors the dashboard: KPI tiles as live links, text tiles as notes, charts re-drawn."""
    name, n = DASHBOARD_SHEET, 2
    while name.lower() in {s.lower() for s in wb.sheetnames}:
        name, n = f"{DASHBOARD_SHEET} ({n})", n + 1
    ws = wb.create_sheet(title=name, index=0)
    wb.active = 0
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    for c in range(2, 2 + 12):
        ws.column_dimensions[get_column_letter(c)].width = DASH_COL_WIDTH
    title = ws.cell(row=1, column=2, value=dash.title or spec.title)
    title.font = Font(name=FONT_NAME, size=18, bold=True, color=ACCENT)
    ws.row_dimensions[1].height = 30

    taken: list[tuple[int, int, int, int]] = []

    def merge(c1: int, r1: int, c2: int, r2: int) -> None:
        if (c1, r1) == (c2, r2):
            return
        box = (c1, r1, c2, r2)
        if any(box[0] <= t[2] and t[0] <= box[2] and box[1] <= t[3] and t[1] <= box[3] for t in taken):
            return
        taken.append(box)
        ws.merge_cells(start_column=c1, start_row=r1, end_column=c2, end_row=r2)

    for tile in sorted(dash.tiles, key=lambda t: (t.y, t.x)):
        x, y = max(0, min(11, tile.x)), max(0, tile.y)
        w, h = max(1, min(12 - x, tile.w)), max(1, min(20, tile.h))
        c1, r1 = 2 + x, DASH_FIRST_ROW + y * DASH_ROWS_PER_UNIT
        c2, r2 = c1 + w - 1, r1 + h * DASH_ROWS_PER_UNIT - 1
        if tile.kind.value == "kpi":
            _dash_kpi(ws, spec, tile, c1, r1, c2, merge)
        elif tile.kind.value == "text":
            head = ws.cell(row=r1, column=c1, value=tile.title or None)
            head.font = BOLD_FONT
            body = ws.cell(row=r1 + 1, column=c1, value=tile.text or None)
            body.alignment = Alignment(wrap_text=True, vertical="top")
            body.font = BODY_FONT
            merge(c1, r1, c2, r1)
            merge(c1, r1 + 1, c2, max(r1 + 1, r2))
        else:
            _dash_chart(wb, ws, spec, tile, c1, r1, w, h)


def _dash_kpi(ws: Worksheet, spec: WorkbookSpec, tile: DashboardTile, c1: int, r1: int, c2: int, merge) -> None:
    label = ws.cell(row=r1, column=c1, value=tile.title or tile.ref or "")
    label.font = KPI_LABEL_FONT
    merge(c1, r1, c2, r1)
    target = parse_cell_ref(tile.ref, spec)
    value = ws.cell(row=r1 + 1, column=c1)
    value.font = KPI_VALUE_FONT
    value.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r1 + 1].height = 30
    merge(c1, r1 + 1, c2, r1 + 1)
    for c in range(c1, c2 + 1):
        for r in (r1, r1 + 1, r1 + 2):
            ws.cell(row=r, column=c).fill = TILE_FILL
    if target is None:
        value.value = "#REF!" if tile.ref else None
        return
    sheet, row, col = target
    value.value = _ref_formula(sheet, row, col)
    value.number_format = _source_format(sheet, row, col)
    compare = parse_cell_ref(tile.compare_ref, spec)
    if compare:
        csheet, crow, ccol = compare
        cref = _ref_formula(csheet, crow, ccol)[1:]
        note = ws.cell(row=r1 + 2, column=c1,
                       value=f'=IFERROR(TEXT({value.value[1:]}/{cref}-1,"+0.0%;-0.0%")&" vs "&TEXT({cref},"#,##0.##"),"")')
        note.font = KPI_NOTE_FONT
        merge(c1, r1 + 2, c2, r1 + 2)


def _dash_chart(wb: Workbook, ws: Worksheet, spec: WorkbookSpec, tile: DashboardTile, c1: int, r1: int, w: int, h: int) -> None:
    sheet = next((s for s in spec.sheets if tile.sheet and s.name.lower() == tile.sheet.lower()), None)
    index = tile.chart or 0
    if sheet is None or not 0 <= index < len(sheet.charts):
        ws.cell(row=r1, column=c1, value=f"{tile.title or 'Chart'}: chart not found").font = KPI_NOTE_FONT
        return
    chart_spec = sheet.charts[index].model_copy(deep=True)
    if tile.title:
        chart_spec.title = tile.title
    chart_spec.anchor = f"{get_column_letter(c1)}{r1}"
    _add_chart(wb[sheet.name], sheet, chart_spec, 0, target=ws,
               size=(w * DASH_COL_WIDTH * _CM_PER_CHAR, h * DASH_ROWS_PER_UNIT * _CM_PER_ROW))
