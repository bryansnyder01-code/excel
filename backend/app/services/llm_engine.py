"""Natural language -> WorkbookSpec, using Claude with structured outputs.

Flow for both "generate" and "edit":
  1. Ask Claude for an `LLMWorkbookResult` (summary + full workbook), constrained to the schema.
  2. Run the formula validator, which auto-corrects safe problems in place.
  3. If real errors remain (bad sheet refs, circular refs, broken syntax), send them back to
     Claude and ask for a corrected workbook, up to `repair_attempts` times.
  4. Return the workbook, the chat summary, and every issue found (fixed or not).

An optional async `progress` callback receives short status lines ("Writing sheet \"Income\"")
while the model streams, so the UI can show what is happening during a long generation.
"""

from __future__ import annotations

import base64
import json
import logging
import re
from collections.abc import Awaitable, Callable

import anthropic

from app.config import Settings, get_settings
from openpyxl.utils import column_index_from_string, get_column_letter

from app.schemas import (
    AICheckResult,
    AIEvalResult,
    AIFillResult,
    ExtractResult,
    ReportSummary,
    AskRequest,
    AskResult,
    ExplainRequest,
    ExplainResult,
    Issue,
    LLMWorkbookResult,
    SheetSpec,
    WorkbookResponse,
    WorkbookSpec,
)
from app.services import prompts
from app.services.formula_validator import has_errors, validate_workbook

log = logging.getLogger(__name__)

# Server-side refusal fallback: if a safety classifier declines, the API retries on a
# suitable model inside the same call instead of returning an empty answer.
FALLBACK_BETA = "server-side-fallback-2026-07-01"

Progress = Callable[[str], Awaitable[None]]

# Sheet objects stream in schema key order, so `"name": "...", "columns"` marks the start of a new sheet.
_SHEET_START = re.compile(r'"name"\s*:\s*"((?:[^"\\]|\\.)*)"\s*,\s*"columns"')
_CHARTS_START = re.compile(r'"charts"\s*:\s*\[\s*\{')
_SCAN_OVERLAP = 512  # characters of already-seen text rescanned for a match split across deltas

# Ask: how much of each sheet is shown to the model.
ASK_MAX_ROWS = 300
ASK_MAX_FORMULAS = 400


class EngineError(RuntimeError):
    """The model could not produce a usable workbook. The message is safe to show users."""


class WorkbookEngine:
    def __init__(self, client: anthropic.AsyncAnthropic | None = None, settings: Settings | None = None):
        self.settings = settings or get_settings()
        self.client = client or anthropic.AsyncAnthropic()

    async def generate(self, prompt: str, progress: Progress | None = None) -> WorkbookResponse:
        return await self._run(prompts.generate_message(prompt), progress)

    async def edit(self, prompt: str, workbook: WorkbookSpec, progress: Progress | None = None) -> WorkbookResponse:
        result = await self._run(prompts.edit_message(prompt, workbook.model_dump_json(exclude_none=True)), progress)
        # Structured outputs can't express free-form dicts, so scenario values come back empty: keep the user's.
        if workbook.scenarios and not any(s.values for s in result.workbook.scenarios):
            result.workbook.scenarios = workbook.scenarios
        _keep_app_state(workbook, result.workbook)
        return result

    # -- AI fill, =AI() cells, photo / PDF extraction, report summaries --------

    async def fill(self, instruction: str, headers: list[str], rows: list[list[str]], output: str) -> list[str]:
        """One short value per row (as text; the route converts numbers). Missing answers come back as ""."""
        message = await self._call(
            [{"role": "user", "content": prompts.fill_message(instruction, headers, rows, output)}],
            system=prompts.FILL_SYSTEM_PROMPT,
            output_format=AIFillResult,
            effort="low",
            max_tokens=min(16000, 1024 + 80 * len(rows)),
        )
        out = [""] * len(rows)
        for v in message.parsed_output.values:
            if 0 <= v.row < len(rows):
                out[v.row] = v.value
        return out

    async def evaluate(self, items: list[dict]) -> dict[str, str]:
        """Results for a batch of =AI(prompt, inputs...) cells, keyed like the request."""
        message = await self._call(
            [{"role": "user", "content": prompts.eval_message(items)}],
            system=prompts.EVAL_SYSTEM_PROMPT,
            output_format=AIEvalResult,
            effort="low",
            max_tokens=min(16000, 1024 + 120 * len(items)),
        )
        keys = {it["key"] for it in items}
        found = {r.key: r.value for r in message.parsed_output.results if r.key in keys}
        return {it["key"]: found.get(it["key"], "") for it in items}

    async def extract(self, data: bytes, media_type: str, hint: str | None = None) -> ExtractResult:
        """Read the tables in an image or PDF (Claude vision)."""
        source = {"type": "base64", "media_type": media_type, "data": base64.b64encode(data).decode()}
        block = {"type": "document" if media_type == "application/pdf" else "image", "source": source}
        message = await self._call(
            [{"role": "user", "content": [block, {"type": "text", "text": prompts.extract_message(hint)}]}],
            system=prompts.EXTRACT_SYSTEM_PROMPT,
            output_format=ExtractResult,
        )
        return message.parsed_output

    async def report_summary(self, title: str, key_numbers: str) -> str:
        message = await self._call(
            [{"role": "user", "content": prompts.report_message(title, key_numbers)}],
            system=prompts.REPORT_SYSTEM_PROMPT,
            output_format=ReportSummary,
            effort="low",
            max_tokens=2000,
        )
        return message.parsed_output.summary.strip()

    async def review(self, workbook: WorkbookSpec, values: dict[str, list[list[str]]], known: str) -> AICheckResult:
        """Look for logic problems (ranges that miss rows, double counting, wrong percentages)."""
        text = prompts.check_message(_ask_tables(workbook, values), _ask_formulas(workbook), known)
        message = await self._call(
            [{"role": "user", "content": text}],
            system=prompts.CHECK_SYSTEM_PROMPT,
            output_format=AICheckResult,
            effort="low",
            max_tokens=8000,
        )
        return message.parsed_output

    async def explain(self, req: ExplainRequest) -> ExplainResult:
        raw = _cell_raw(req.workbook, req.sheet, req.cell)
        text = prompts.explain_message(
            req.workbook.model_dump_json(exclude_none=True), req.sheet, req.cell, raw, req.computed, req.question
        )
        message = await self._call(
            [{"role": "user", "content": text}],
            system=prompts.EXPLAIN_SYSTEM_PROMPT,
            output_format=ExplainResult,
            effort="low",
            max_tokens=8000,
        )
        return message.parsed_output

    async def ask(self, req: AskRequest) -> AskResult:
        """Answer a question about the workbook's data from the values the grid currently shows."""
        text = prompts.ask_message(_ask_tables(req.workbook, req.values), _ask_formulas(req.workbook), req.question)
        message = await self._call(
            [{"role": "user", "content": text}],
            system=prompts.ASK_SYSTEM_PROMPT,
            output_format=AskResult,
            effort="low",
            max_tokens=8000,
        )
        result: AskResult = message.parsed_output
        result.references = result.references[:5]
        return result

    async def _run(self, user_text: str, progress: Progress | None = None) -> WorkbookResponse:
        report = _Reporter(progress)
        await report("Reading your request")
        messages: list[dict] = [{"role": "user", "content": user_text}]
        attempts = self.settings.repair_attempts + 1

        for attempt in range(attempts):
            message = await self._call(messages, progress=report if progress else None)
            result: LLMWorkbookResult = message.parsed_output
            await report("Checking formulas")
            issues = validate_workbook(result.workbook)

            if not has_errors(issues) or attempt == attempts - 1:
                return WorkbookResponse(workbook=result.workbook, summary=result.summary, issues=issues)

            n = sum(1 for i in issues if i.severity == "error" and not i.fixed)
            await report(f"Fixing {n} formula {'problem' if n == 1 else 'problems'}")
            log.info("Workbook had %d errors, asking the model to repair (attempt %d)", len(issues), attempt + 1)
            # Append-only history: the model's full turn (thinking blocks included), then the errors.
            messages.append({"role": "assistant", "content": message.content})
            messages.append({"role": "user", "content": prompts.repair_message(_format_issues(issues))})

        raise AssertionError("unreachable")

    async def _call(self, messages: list[dict], system: str = prompts.SYSTEM_PROMPT, output_format=LLMWorkbookResult,
                    effort: str | None = None, max_tokens: int | None = None, progress: Progress | None = None):
        try:
            async with self.client.messages.stream(
                model=self.settings.model,
                max_tokens=max_tokens or self.settings.max_tokens,
                system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral"}}],
                messages=messages,
                output_format=output_format,
                output_config={"effort": effort or self.settings.effort},
                extra_headers={"anthropic-beta": FALLBACK_BETA},
                extra_body={"fallbacks": "default"},
            ) as stream:
                if progress:
                    await _watch_stream(stream, progress)
                message = await stream.get_final_message()
        except anthropic.RateLimitError as e:
            raise EngineError("The AI service is busy right now. Please try again in a moment.") from e
        except anthropic.BadRequestError as e:
            log.exception("Bad request to the model")
            raise EngineError("The request could not be processed. Try rephrasing it.") from e
        except (anthropic.AuthenticationError, anthropic.PermissionDeniedError) as e:
            log.exception("API key rejected")
            hint = " Change it from the app's menu (Change API key)." if self.settings.static_dir else ""
            raise EngineError(f"Anthropic didn't accept the API key.{hint}") from e
        except anthropic.APIConnectionError as e:
            log.exception("Couldn't reach the model")
            raise EngineError("Can't reach the AI service. Check your internet connection and try again.") from e
        except anthropic.APIStatusError as e:
            log.exception("Model call failed")
            raise EngineError("The AI service is unavailable. Please try again.") from e

        if message.stop_reason == "refusal":
            raise EngineError("That request can't be turned into a spreadsheet. Try describing it differently.")
        if message.stop_reason == "max_tokens":
            raise EngineError("That workbook is too large to generate in one go. Try asking for fewer sheets or rows.")
        if message.parsed_output is None:
            raise EngineError("The AI returned an unreadable answer. Please try again.")
        return message


class _Reporter:
    """Wraps the caller's progress callback: drops duplicate lines and is a no-op without a callback."""

    def __init__(self, callback: Progress | None):
        self.callback = callback
        self.sent: set[str] = set()

    async def __call__(self, text: str) -> None:
        if self.callback and text not in self.sent:
            self.sent.add(text)
            await self.callback(text)


async def _watch_stream(stream, progress: Progress) -> None:
    """Turn raw stream events into progress lines while the model writes its JSON."""
    text, scan_from = "", 0
    async for event in stream:
        if event.type == "content_block_start" and event.content_block.type in ("thinking", "redacted_thinking"):
            await progress("Planning the layout")
        elif event.type == "content_block_delta" and event.delta.type == "text_delta":
            text += event.delta.text
            for m in _SHEET_START.finditer(text, scan_from):
                await progress(f'Writing sheet "{_unescape(m.group(1))}"')
            if _CHARTS_START.search(text, scan_from):
                await progress("Adding charts")
            scan_from = max(0, len(text) - _SCAN_OVERLAP)


def _unescape(raw: str) -> str:
    try:
        return json.loads(f'"{raw}"')
    except ValueError:
        return raw


def _sheet_grid(sheet: SheetSpec) -> list[list[str]]:
    """Fallback when the client sent no values for a sheet: header plus raw cell contents."""
    grid = [[c.header for c in sheet.columns]]
    for row in sheet.rows:
        grid.append(["" if c.value is None else str(c.value) for c in row.cells])
    return grid


def _ask_tables(wb: WorkbookSpec, values: dict[str, list[list[str]]]) -> str:
    """Each sheet as a pipe table with Excel column letters and row numbers."""
    parts = []
    for sheet in wb.sheets:
        grid = values.get(sheet.name)
        if grid is None:
            grid = _sheet_grid(sheet)
        shown = grid[:ASK_MAX_ROWS]
        width = max((len(r) for r in shown), default=0)
        lines = ["    | " + " | ".join(get_column_letter(c + 1) for c in range(width))]
        for r, row in enumerate(shown, start=1):
            cells = [str(v).replace("|", "/").replace("\n", " ") for v in row]
            lines.append(f"{r:>3} | " + " | ".join(cells))
        note = f" (showing the first {ASK_MAX_ROWS} of {len(grid)} rows)" if len(grid) > ASK_MAX_ROWS else ""
        parts.append(f'<sheet name="{sheet.name}"{note}>\n' + "\n".join(lines) + "\n</sheet>")
    return "\n\n".join(parts)


def _ask_formulas(wb: WorkbookSpec) -> str:
    """Formula cells, one per line, so the model can explain how numbers are computed."""
    lines = []
    for sheet in wb.sheets:
        found = [
            f"'{sheet.name}'!{get_column_letter(c)}{r} {cell.value}"
            for r, row in enumerate(sheet.rows, start=2)
            for c, cell in enumerate(row.cells, start=1)
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        lines += found[:ASK_MAX_FORMULAS]
        if len(found) > ASK_MAX_FORMULAS:
            lines.append(f"'{sheet.name}': {len(found) - ASK_MAX_FORMULAS} more formulas omitted")
    return "\n".join(lines) or "(no formulas)"


def _format_issues(issues: list[Issue]) -> str:
    lines = []
    for i in issues:
        if i.severity != "error" or i.fixed:
            continue
        where = " ".join(p for p in (f"sheet '{i.sheet}'" if i.sheet else "", f"cell {i.cell}" if i.cell else "") if p)
        lines.append(f"- {where}: {i.message}" if where else f"- {i.message}")
    return "\n".join(lines)


def _cell_raw(wb: WorkbookSpec, sheet_name: str, a1: str) -> str | None:
    sheet = next((s for s in wb.sheets if s.name.lower() == sheet_name.lower()), None)
    letters = "".join(ch for ch in a1 if ch.isalpha()).upper()
    digits = "".join(ch for ch in a1 if ch.isdigit())
    if sheet is None or not letters or not digits:
        return None
    col, row = column_index_from_string(letters) - 1, int(digits)
    if row == 1:
        return sheet.columns[col].header if col < len(sheet.columns) else None
    if row - 2 < len(sheet.rows) and col < len(sheet.rows[row - 2].cells):
        v = sheet.rows[row - 2].cells[col].value
        return None if v is None else str(v)
    return None


def _keep_app_state(before: WorkbookSpec, after: WorkbookSpec) -> None:
    """Things the model has no reason to touch (dashboard, data connections, protection, =AI() results) survive an
    edit even when the model's answer leaves them out."""
    if before.dashboard and after.dashboard is None:
        after.dashboard = before.dashboard
    if before.connections and not after.connections:
        after.connections = before.connections
    old_sheets = {s.name: s for s in before.sheets}
    for sheet in after.sheets:
        old = old_sheets.get(sheet.name)
        if old is None:
            continue
        if old.locked and not sheet.locked:
            sheet.locked = True
        if old.locked_ranges and not sheet.locked_ranges:
            sheet.locked_ranges = list(old.locked_ranges)
        for r, row in enumerate(sheet.rows):
            for c, cell in enumerate(row.cells):
                if cell.ai_value is None and r < len(old.rows) and c < len(old.rows[r].cells):
                    prev = old.rows[r].cells[c]
                    if prev.ai_value is not None and prev.value == cell.value:
                        cell.ai_value = prev.ai_value
