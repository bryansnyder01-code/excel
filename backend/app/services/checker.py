""""Check my workbook": deterministic spreadsheet lint plus an optional AI logic review.

Deterministic checks (no AI, always run):
  * cells that display an Excel error (#DIV/0!, #REF!, ...)                    -> error
  * total / subtotal cells with a typed number next to formula totals          -> warning, fix = SUM above
  * a formula whose shape differs from the rest of its column                  -> warning, fix = majority formula
  * a typed number in a column that is mostly formulas                         -> warning, fix = majority formula
  * a formula that points at an empty cell inside the sheet's used range       -> info
  * the formula validator's own (unfixed) issues

"Shape" is the R1C1 form of a formula: relative references become row/column offsets from the
cell itself, so =B2*C2 in row 2 and =B9*C9 in row 9 have the same shape.

With `use_ai`, one low-effort Claude call looks for logic problems the rules can't see (ranges
that miss rows, double-counted totals, wrong percentages). Its findings are merged in; if the
call fails the deterministic findings are returned on their own.
"""

from __future__ import annotations

import hashlib
import logging
import re
from collections import Counter

from openpyxl.utils import column_index_from_string, get_column_letter

from app.schemas import CheckRequest, CheckResponse, Finding, RowKind, SheetSpec, WorkbookSpec
from app.services.formula_validator import STRING_RE, iter_refs, validate_workbook
from app.services.llm_engine import EngineError

log = logging.getLogger(__name__)

MAX_FINDINGS = 50
MAX_AI_FINDINGS = 8
SEVERITY_RANK = {"error": 0, "warning": 1, "info": 2}

EXCEL_ERRORS = {
    "#DIV/0!": "divides by zero or by an empty cell",
    "#REF!": "points at a cell that no longer exists",
    "#NAME?": "uses a function or name Excel doesn't recognise",
    "#VALUE!": "mixes text into a calculation that needs numbers",
    "#N/A": "looks up a value that wasn't found",
    "#NUM!": "produces a number that is invalid or too large",
    "#CYCLE!": "depends on its own result (a circular reference)",
}

# A cell reference, with an optional sheet prefix kept verbatim so 'Q1 Data'!B2 isn't read as cell Q1.
_CELL_TOKEN = re.compile(
    r"(?<![A-Za-z0-9_\.])(?P<prefix>'(?:[^']|'')+'!|[A-Za-z_][\w\.]*!)?"
    r"(?P<cabs>\$?)(?P<col>[A-Z]{1,3})(?P<rabs>\$?)(?P<row>[0-9]+)(?![A-Za-z0-9_\(!])"
)
_A1 = re.compile(r"^\$?([A-Za-z]{1,3})\$?([0-9]+)$")


def finding_id(kind: str, sheet: str, cell: str) -> str:
    return hashlib.sha1(f"{kind}|{sheet}|{cell}".encode()).hexdigest()[:16]


def _finding(severity: str, sheet: str, cell: str, kind: str, message: str, fix: str | None = None) -> Finding:
    return Finding(id=finding_id(kind, sheet, cell), severity=severity, sheet=sheet, cell=cell, kind=kind,
                   message=message, fix=fix)


def _is_formula(value) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _map_refs(formula: str, fn) -> str:
    """Apply `fn(match) -> str` to every cell reference outside string literals."""
    out, pos = [], 0
    for m in STRING_RE.finditer(formula):
        out.append(_CELL_TOKEN.sub(fn, formula[pos:m.start()]))
        out.append(m.group(0))
        pos = m.end()
    out.append(_CELL_TOKEN.sub(fn, formula[pos:]))
    return "".join(out)


def formula_shape(formula: str, row: int, col: int) -> str:
    """R1C1-style shape: relative parts become offsets from (row, col), absolute parts stay fixed."""

    def r1c1(m: re.Match) -> str:
        c, r = column_index_from_string(m.group("col")), int(m.group("row"))
        rp = f"R{r}" if m.group("rabs") else f"R[{r - row}]"
        cp = f"C{c}" if m.group("cabs") else f"C[{c - col}]"
        return (m.group("prefix") or "") + rp + cp

    return _map_refs(formula.replace(" ", ""), r1c1)


def move_formula(formula: str, from_row: int, to_row: int) -> str:
    """The formula as if copied down (or up) from `from_row` to `to_row`; absolute rows stay put."""
    delta = to_row - from_row

    def shift(m: re.Match) -> str:
        r = int(m.group("row"))
        if not m.group("rabs"):
            r = max(1, r + delta)
        return f"{m.group('prefix') or ''}{m.group('cabs')}{m.group('col')}{m.group('rabs')}{r}"

    return _map_refs(formula, shift)


def _addr(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


# ---------------------------------------------------------------------------
# Deterministic rules
# ---------------------------------------------------------------------------


def _error_cells(sheet: SheetSpec, grid: list[list[str]] | None) -> list[Finding]:
    out = []
    for r, row in enumerate(grid or []):
        if r == 0:
            continue
        for c, text in enumerate(row):
            shown = str(text).strip().upper() if text is not None else ""
            if shown in EXCEL_ERRORS:
                out.append(_finding("error", sheet.name, _addr(c + 1, r + 1), "excel_error",
                                    f"This cell shows {shown}: its formula {EXCEL_ERRORS[shown]}."))
    return out


def _hardcoded_totals(sheet: SheetSpec) -> list[Finding]:
    out = []
    block_start = 2  # first data row of the current block
    for r_idx, row in enumerate(sheet.rows):
        excel_row = r_idx + 2
        if row.kind in (RowKind.total, RowKind.subtotal):
            if any(_is_formula(c.value) for c in row.cells):
                label = "subtotal" if row.kind == RowKind.subtotal else "total"
                for c_idx, cell in enumerate(row.cells):
                    if not _is_number(cell.value):
                        continue
                    letter = get_column_letter(c_idx + 1)
                    end = excel_row - 1
                    fix = f"=SUM({letter}{block_start}:{letter}{end})" if end >= block_start else None
                    out.append(_finding(
                        "warning", sheet.name, f"{letter}{excel_row}", "hardcoded_total",
                        f"This {label} is a typed number, while the other {label}s in the row are formulas, "
                        "so it won't update when the data changes.", fix))
        if row.kind != RowKind.data:
            block_start = excel_row + 1
    return out


def _column_formulas(sheet: SheetSpec) -> list[Finding]:
    """Inconsistent formulas and typed numbers in formula columns, per column, data rows only."""
    out = []
    ncols = max([len(sheet.columns)] + [len(r.cells) for r in sheet.rows])
    for c_idx in range(ncols):
        col = c_idx + 1
        formulas: list[tuple[int, str, str]] = []  # (excel row, formula, shape)
        numbers: list[int] = []
        filled = 0
        for r_idx, row in enumerate(sheet.rows):
            if row.kind != RowKind.data or c_idx >= len(row.cells):
                continue
            value = row.cells[c_idx].value
            if value is None or value == "":
                continue
            filled += 1
            if _is_formula(value):
                formulas.append((r_idx + 2, value, formula_shape(value, r_idx + 2, col)))
            elif _is_number(value):
                numbers.append(r_idx + 2)
        if len(formulas) < 2:
            continue
        shape, count = Counter(s for _, _, s in formulas).most_common(1)[0]
        majority = count >= 2 and count * 2 > len(formulas)
        model_row, model = next(((r, f) for r, f, s in formulas if s == shape), (0, ""))
        letter = get_column_letter(col)
        if majority:
            for r, f, s in formulas:
                if s != shape:
                    out.append(_finding(
                        "warning", sheet.name, f"{letter}{r}", "inconsistent_formula",
                        f"This formula ({f}) doesn't follow the pattern of the rest of column {letter}.",
                        move_formula(model, model_row, r)))
        if majority and len(formulas) * 2 > filled:
            for r in numbers:
                out.append(_finding(
                    "warning", sheet.name, f"{letter}{r}", "number_in_formula_column",
                    f"This is a typed number, but the rest of column {letter} is calculated with formulas.",
                    move_formula(model, model_row, r)))
    return out


def _empty_refs(wb: WorkbookSpec) -> list[Finding]:
    sheets = {s.name.lower(): s for s in wb.sheets}
    out = []
    for sheet in wb.sheets:
        for r_idx, row in enumerate(sheet.rows):
            for c_idx, cell in enumerate(row.cells):
                if not _is_formula(cell.value):
                    continue
                empty = []
                for ref in iter_refs(cell.value):
                    target = sheets.get((ref.sheet or sheet.name).lower())
                    if target is None or ref.min_row is None or (ref.min_row, ref.min_col) != (ref.max_row, ref.max_col):
                        continue
                    r, c = ref.min_row, ref.min_col
                    if r < 2 or r - 2 >= len(target.rows) or c > len(target.columns):
                        continue  # header row or outside the used range
                    cells = target.rows[r - 2].cells
                    value = cells[c - 1].value if c - 1 < len(cells) else None
                    if value is None or value == "":
                        where = "" if target is sheet else f"'{target.name}'!"
                        empty.append(f"{where}{_addr(c, r)}")
                if empty:
                    names = ", ".join(dict.fromkeys(empty))
                    out.append(_finding("info", sheet.name, _addr(c_idx + 1, r_idx + 2), "references_empty",
                                        f"This formula uses {names}, which {'is' if len(empty) == 1 else 'are'} empty."))
    return out


def _validator_findings(wb: WorkbookSpec) -> list[Finding]:
    copy = wb.model_copy(deep=True)
    out = []
    for issue in validate_workbook(copy, autocorrect=False):
        if issue.fixed or not issue.sheet:
            continue
        cell = (issue.cell or "").split(":")[0].strip()
        m = _A1.match(cell)
        cell = f"{m.group(1).upper()}{m.group(2)}" if m else "A1"
        out.append(_finding(issue.severity if issue.severity in SEVERITY_RANK else "warning", issue.sheet, cell,
                            "validation", issue.message))
    return out


def deterministic_findings(wb: WorkbookSpec, values: dict[str, list[list[str]]]) -> list[Finding]:
    out: list[Finding] = []
    for sheet in wb.sheets:
        out += _error_cells(sheet, values.get(sheet.name))
        out += _hardcoded_totals(sheet)
        out += _column_formulas(sheet)
    out += _empty_refs(wb)
    out += _validator_findings(wb)
    return out


# ---------------------------------------------------------------------------
# Merging and the endpoint entry point
# ---------------------------------------------------------------------------


def normalize_ai_findings(wb: WorkbookSpec, raw) -> list[Finding]:
    """Turn the model's findings into Findings on real sheets and cells; drop anything unusable."""
    names = {s.name.lower(): s.name for s in wb.sheets}
    out = []
    for f in raw[:MAX_AI_FINDINGS]:
        sheet, cell = f.sheet.strip(), f.cell.strip()
        if "!" in cell:
            prefix, _, cell = cell.rpartition("!")
            sheet = prefix.strip("'") or sheet
        sheet = names.get(sheet.strip("'").lower())
        m = _A1.match(cell.split(":")[0])
        if not sheet or not m or not f.message.strip():
            continue
        cell = f"{m.group(1).upper()}{m.group(2)}"
        severity = f.severity.strip().lower()
        severity = severity if severity in SEVERITY_RANK else "warning"
        kind = re.sub(r"[^a-z0-9_]+", "_", (f.kind or "logic").strip().lower()).strip("_") or "logic"
        fix = f.fix.strip() if f.fix and f.fix.strip() else None
        out.append(_finding(severity, sheet, cell, kind, f.message.strip(), fix))
    return out


def merge_findings(*groups: list[Finding]) -> list[Finding]:
    """Errors first; one finding per (sheet, cell), the most severe (earlier groups win ties); at most 50."""
    ordered = sorted(
        ((SEVERITY_RANK.get(f.severity, 2), g, i, f) for g, group in enumerate(groups) for i, f in enumerate(group)),
        key=lambda t: t[:3],
    )
    seen: set[tuple[str, str]] = set()
    out = []
    for *_, f in ordered:
        key = (f.sheet.lower(), f.cell.upper())
        if key in seen:
            continue
        seen.add(key)
        out.append(f)
    return out[:MAX_FINDINGS]


def summarize(findings: list[Finding]) -> str:
    counts = Counter(f.severity for f in findings)
    parts = []
    for sev, one, many in (("error", "error", "errors"), ("warning", "warning", "warnings"), ("info", "note", "notes")):
        if counts[sev]:
            parts.append(f"{counts[sev]} {one if counts[sev] == 1 else many}")
    if not parts:
        return "No problems found."
    return "Found " + (parts[0] if len(parts) == 1 else ", ".join(parts[:-1]) + " and " + parts[-1]) + "."


def known_findings_text(findings: list[Finding]) -> str:
    return "\n".join(f"- '{f.sheet}'!{f.cell}: {f.message}" for f in findings) or "(none)"


async def check_workbook(req: CheckRequest, engine=None) -> CheckResponse:
    """Run the checks; `engine` (a WorkbookEngine) is only used when `req.use_ai` is set."""

    base = deterministic_findings(req.workbook, req.values)
    ai: list[Finding] = []
    note = ""
    if req.use_ai and engine is not None:
        try:
            result = await engine.review(req.workbook, req.values, known_findings_text(base))
            ai = normalize_ai_findings(req.workbook, result.findings)
        except EngineError as e:
            log.info("AI review unavailable: %s", e)
            note = " (the AI review was unavailable, so only the automatic checks ran)"
    findings = merge_findings(base, ai)
    return CheckResponse(summary=summarize(findings)[:-1] + note + ".", findings=findings)
