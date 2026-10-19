"""Build a pivot-style summary sheet from a data sheet, using live SUMIFS-family formulas.

Unlike an Excel PivotTable, the result is ordinary cells and formulas, so it previews in the
grid, recalculates when the source data changes, and survives every export format.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from app.schemas import (
    Aggregation,
    CellSpec,
    ChartSeriesSpec,
    ChartSpec,
    ChartType,
    ColumnSpec,
    NumberFormat,
    PivotRequest,
    RowKind,
    RowSpec,
    SheetSpec,
    WorkbookSpec,
)
from app.services.formula_validator import sanitize_sheet_name

MAX_GROUPS = 200
MAX_COLUMN_GROUPS = 24

_PLAIN = {Aggregation.sum: "SUM", Aggregation.average: "AVERAGE", Aggregation.count: "COUNTA",
          Aggregation.min: "MIN", Aggregation.max: "MAX"}
_IFS = {Aggregation.sum: "SUMIFS", Aggregation.average: "AVERAGEIFS", Aggregation.count: "COUNTIFS",
        Aggregation.min: "MINIFS", Aggregation.max: "MAXIFS"}
_LABEL = {Aggregation.sum: "Total", Aggregation.average: "Average", Aggregation.count: "Count",
          Aggregation.min: "Min", Aggregation.max: "Max"}


class PivotError(ValueError):
    """The request doesn't fit the data. The message is safe to show users."""


def _find_column(sheet: SheetSpec, header: str) -> int:
    for i, col in enumerate(sheet.columns):
        if col.header.strip().lower() == header.strip().lower():
            return i
    raise PivotError(f'"{sheet.name}" has no column called "{header}".')


def _distinct(sheet: SheetSpec, col: int, limit: int, field: str) -> list:
    seen: dict = {}
    for row in sheet.rows:
        if row.kind != RowKind.data or col >= len(row.cells):
            continue
        v = row.cells[col].value
        if v is None or v == "" or (isinstance(v, str) and v.startswith("=")):
            continue
        key = v.strip().lower() if isinstance(v, str) else v
        seen.setdefault(key, v.strip() if isinstance(v, str) else v)
    values = list(seen.values())
    if not values:
        raise PivotError(f'The "{field}" column has no values to group by.')
    if len(values) > limit:
        raise PivotError(f'"{field}" has {len(values)} different values; summaries support up to {limit}.')
    try:
        return sorted(values, key=lambda v: (isinstance(v, str), v))
    except TypeError:
        return values


def build_pivot(wb: WorkbookSpec, req: PivotRequest) -> tuple[WorkbookSpec, str]:
    src = next((s for s in wb.sheets if s.name.lower() == req.source_sheet.lower()), None)
    if src is None:
        raise PivotError(f'There is no sheet called "{req.source_sheet}".')
    agg = req.aggregation
    r_idx = _find_column(src, req.row_field)
    c_idx = _find_column(src, req.column_field) if req.column_field else None
    if agg != Aggregation.count and not req.value_field:
        raise PivotError("Choose which column of numbers to summarise.")
    v_idx = _find_column(src, req.value_field) if req.value_field else r_idx

    last = 1 + max((i + 1 for i, r in enumerate(src.rows) if r.kind == RowKind.data), default=0)
    if last < 2:
        raise PivotError(f'"{src.name}" has no data rows.')
    q = "'" + src.name.replace("'", "''") + "'!"

    def rng(i: int) -> str:
        L = get_column_letter(i + 1)
        return f"{q}${L}$2:${L}${last}"

    value_rng, row_rng = rng(v_idx), rng(r_idx)
    col_rng = rng(c_idx) if c_idx is not None else None
    ifs, plain = _IFS[agg], _PLAIN[agg]

    def formula(row_crit: str | None, col_crit: str | None) -> str:
        crit = []
        if row_crit:
            crit += [row_rng, row_crit]
        if col_crit:
            crit += [col_rng, col_crit]
        if not crit:
            return f"={plain}({value_rng})"
        if agg == Aggregation.count:
            return f"={ifs}({','.join(crit)})"
        f = f"{ifs}({value_rng},{','.join(crit)})"
        return f"=IFERROR({f},0)" if agg in (Aggregation.average, Aggregation.min, Aggregation.max) else f"={f}"

    groups = _distinct(src, r_idx, MAX_GROUPS, req.row_field)
    col_groups = _distinct(src, c_idx, MAX_COLUMN_GROUPS, req.column_field) if c_idx is not None else []

    if agg == Aggregation.count:
        fmt = NumberFormat.integer
    else:
        fmt = src.columns[v_idx].format
        if fmt in (NumberFormat.text, NumberFormat.date):
            fmt = NumberFormat.decimal if agg == Aggregation.average else NumberFormat.integer
        if agg == Aggregation.average and fmt == NumberFormat.integer:
            fmt = NumberFormat.decimal

    value_label = f"{_LABEL[agg]} of {req.value_field}" if agg != Aggregation.count else "Count"
    columns = [ColumnSpec(header=src.columns[r_idx].header, format=src.columns[r_idx].format)]
    if col_groups:
        columns += [ColumnSpec(header=str(g if not (isinstance(g, float) and g.is_integer()) else int(g)), format=fmt)
                    for g in col_groups]
        columns.append(ColumnSpec(header=_LABEL[agg] if agg != Aggregation.count else "Total", format=fmt))
    else:
        columns.append(ColumnSpec(header=value_label, format=fmt))

    n_val_cols = len(columns) - 1
    rows: list[RowSpec] = []
    for i, g in enumerate(groups):
        excel_row = i + 2
        cells = [CellSpec(value=g)]
        for j in range(n_val_cols):
            letter = get_column_letter(j + 2)
            is_total_col = col_groups and j == n_val_cols - 1
            col_crit = f"{letter}$1" if col_groups and not is_total_col else None
            cells.append(CellSpec(value=formula(f"$A{excel_row}", col_crit)))
        rows.append(RowSpec(cells=cells))

    total = [CellSpec(value="Grand total" if agg in (Aggregation.sum, Aggregation.count) else f"All ({_LABEL[agg].lower()})")]
    for j in range(n_val_cols):
        letter = get_column_letter(j + 2)
        is_total_col = col_groups and j == n_val_cols - 1
        total.append(CellSpec(value=formula(None, f"{letter}$1" if col_groups and not is_total_col else None)))
    rows.append(RowSpec(kind=RowKind.total, cells=total))

    taken = {s.name.lower() for s in wb.sheets}
    base = req.name or (f"{req.value_field or 'Count'} by {req.row_field}" + (f" & {req.column_field}" if req.column_field else ""))
    name = sanitize_sheet_name(base, taken)

    charts = []
    if len(groups) <= 15:
        last_row = len(groups) + 1
        series_cols = range(2, len(columns)) if col_groups else [2]
        charts.append(ChartSpec(
            type=ChartType.column,
            title=f"{value_label} by {req.row_field}" + (f" and {req.column_field}" if req.column_field else ""),
            categories=f"A2:A{last_row}",
            series=[ChartSeriesSpec(name=columns[c - 1].header if col_groups else value_label,
                                    range=f"{get_column_letter(c)}2:{get_column_letter(c)}{last_row}") for c in series_cols],
        ))

    out = wb.model_copy(deep=True)
    out.sheets.append(SheetSpec(name=name, columns=columns, rows=rows, charts=charts))
    summary = (f'Added a "{name}" tab with {value_label.lower()} for each {req.row_field}'
               + (f" split by {req.column_field}" if req.column_field else "")
               + f", using live formulas that update when {src.name} changes.")
    return out, summary
