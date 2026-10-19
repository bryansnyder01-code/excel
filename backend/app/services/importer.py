"""Turn an uploaded .xlsx or .csv into a WorkbookSpec so it can be previewed, edited by chat and re-exported.

Formulas are kept as formulas. The first row of each sheet becomes the header row.
Basic styling (bold/italic/underline, explicit font and fill colours, alignment,
wrap text), row heights, merged cells, number formats, autofilter selections and
dropdown lists are carried over; charts and pivot tables are not.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from app.schemas import (
    CellSpec,
    ColumnSpec,
    DataValidationSpec,
    FilterSpec,
    Issue,
    NumberFormat,
    RowKind,
    RowSpec,
    SheetSpec,
    ValidationType,
    WorkbookSpec,
)

MAX_ROWS = 20000
# Fills SheetCraft's own theme applies; re-importing an exported file shouldn't freeze them as overrides.
_THEME_FILLS = {"#FFFFFF", "#F7F9FC", "#DDEBF7", "#BDD7EE"}
MAX_COLS = 100
# Header height SheetCraft's builder always writes; not worth keeping as an explicit setting.
_DEFAULT_HEADER_HEIGHT = 22


class ImportError_(ValueError):
    """The file could not be read. The message is safe to show users."""


def import_file(filename: str, data: bytes) -> tuple[WorkbookSpec, list[Issue]]:
    name = filename.lower()
    title = re.sub(r"\.(xlsx|xlsm|csv|tsv|txt)$", "", filename, flags=re.I) or "Imported workbook"
    if name.endswith((".xlsx", ".xlsm")):
        return _import_xlsx(title, data)
    if name.endswith((".csv", ".tsv", ".txt")):
        return _import_csv(title, data)
    raise ImportError_("Please upload an .xlsx or .csv file.")


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------


def _import_xlsx(title: str, data: bytes) -> tuple[WorkbookSpec, list[Issue]]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=False)
    except Exception as e:  # openpyxl raises many types for corrupt files
        raise ImportError_("That file couldn't be opened as an Excel workbook.") from e

    issues: list[Issue] = []
    sheets: list[SheetSpec] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        max_row, max_col = ws.max_row, ws.max_column
        if max_row > MAX_ROWS + 1 or max_col > MAX_COLS:
            issues.append(Issue(severity="warning", sheet=ws.title,
                                message=f"Only the first {MAX_ROWS} rows and {MAX_COLS} columns were imported."))
        max_row, max_col = min(max_row, MAX_ROWS + 1), min(max_col, MAX_COLS)
        if ws._charts:
            issues.append(Issue(severity="warning", sheet=ws.title,
                                message="Charts in this sheet were not imported; ask the AI to recreate them."))

        header_cells = [ws.cell(row=1, column=c) for c in range(1, max_col + 1)]
        columns = [
            ColumnSpec(header=str(h.value) if h.value is not None else f"Column {get_column_letter(i + 1)}")
            for i, h in enumerate(header_cells)
        ]
        rows: list[RowSpec] = []
        fmt_votes: list[dict[NumberFormat, int]] = [{} for _ in columns]
        # (format, currency, decimals) votes, so the column can carry the common currency / precision.
        detail_votes: list[dict[tuple, int]] = [{} for _ in columns]
        for r in range(2, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                spec, fmt = _cell_spec(cell)
                cells.append(spec)
                if fmt and spec.value is not None:
                    fmt_votes[c - 1][fmt] = fmt_votes[c - 1].get(fmt, 0) + 1
                    key = (fmt, spec.currency, spec.decimals)
                    detail_votes[c - 1][key] = detail_votes[c - 1].get(key, 0) + 1
            while cells and cells[-1].value is None and not _styled(cells[-1]):
                cells.pop()
            rows.append(RowSpec(kind=_row_kind(cells), cells=cells, height=_row_height(ws, r)))
        while rows and not rows[-1].cells and rows[-1].height is None:
            rows.pop()

        # Column format = most common cell format; cells matching it drop their override.
        for i, votes in enumerate(fmt_votes):
            if votes:
                columns[i].format = max(votes, key=votes.get)
            same = {k: n for k, n in detail_votes[i].items() if k[0] == columns[i].format}
            if same:
                _, columns[i].currency, columns[i].decimals = max(same, key=same.get)
            for row in rows:
                if i < len(row.cells) and row.cells[i].format == columns[i].format:
                    row.cells[i].format = None
                    if (row.cells[i].currency, row.cells[i].decimals) == (columns[i].currency, columns[i].decimals):
                        row.cells[i].currency = row.cells[i].decimals = None
            width = ws.column_dimensions[get_column_letter(i + 1)].width
            if width and ws.column_dimensions[get_column_letter(i + 1)].customWidth:
                columns[i].width = round(width, 1)

        freeze_header, freeze_columns = _freeze(ws.freeze_panes)
        header_height = _row_height(ws, 1)
        locked, locked_ranges = _protection(ws, max_row, max_col)
        sheets.append(SheetSpec(
            name=ws.title, columns=columns, rows=rows,
            freeze_header=freeze_header, freeze_columns=freeze_columns,
            show_gridlines=ws.sheet_view.showGridLines is not False,
            validations=_validations(ws),
            merges=_merges(ws, max_row, max_col),
            filters=_filters(ws, max_col),
            header_height=header_height if header_height != _DEFAULT_HEADER_HEIGHT else None,
            locked=locked,
            locked_ranges=locked_ranges,
        ))

    if not sheets:
        raise ImportError_("That workbook has no visible sheets.")
    return WorkbookSpec(title=wb.properties.title or title, sheets=sheets), issues


def _cell_spec(cell: Cell) -> tuple[CellSpec, NumberFormat | None]:
    if isinstance(cell, MergedCell):  # covered by a merge: the top-left cell holds the value and style
        return CellSpec(), None
    value = cell.value
    if isinstance(value, datetime):
        value = value.date().isoformat() if value.time() == time() else value.isoformat(sep=" ")
    elif isinstance(value, date):
        value = value.isoformat()
    elif isinstance(value, time):
        value = value.isoformat()
    elif isinstance(value, int) and not isinstance(value, bool):
        value = float(value)
    elif value is not None and not isinstance(value, (str, float, bool)):
        value = str(value)  # ArrayFormula, rich text, etc.
        if hasattr(cell.value, "text") and cell.value.text:
            value = cell.value.text

    fmt = _number_format(cell)
    currency, decimals = _currency_and_decimals(cell.number_format or "", fmt)
    font, fill = cell.font, cell.fill
    spec = CellSpec(
        value=value,
        format=fmt,
        bold=bool(font and font.b),
        italic=bool(font and font.i),
        underline=bool(font and font.u),
        font_color=_rgb(font.color) if font else None,
        fill_color=_rgb(fill.fgColor) if fill is not None and fill.fill_type == "solid" else None,
        align=cell.alignment.horizontal if cell.alignment and cell.alignment.horizontal in ("left", "center", "right") else None,
        note=cell.comment.text if cell.comment else None,
        currency=currency,
        decimals=decimals,
        wrap=bool(cell.alignment and cell.alignment.wrap_text),
    )
    if spec.font_color in ("#000000",):
        spec.font_color = None
    if spec.fill_color in _THEME_FILLS:
        spec.fill_color = None
    return spec, fmt


def _rgb(color) -> str | None:
    """Only explicit ARGB colours carry over; theme/indexed colours fall back to the default theme."""
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str) or len(color.rgb) != 8:
        return None
    return "#" + color.rgb[2:].upper()


def _number_format(cell: Cell) -> NumberFormat | None:
    nf = (cell.number_format or "General").lower()
    v = cell.value
    if cell.is_date or isinstance(v, (date, datetime)):
        return NumberFormat.date
    if "%" in nf:
        return NumberFormat.percent
    if any(sym in nf for sym in ("$", "€", "£", "¥", "[$")) or "accounting" in nf:
        return NumberFormat.currency
    if nf == "@":
        return NumberFormat.text
    if re.search(r"0\.0", nf):
        return NumberFormat.decimal
    if "0" in nf or "#" in nf:
        return NumberFormat.integer
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return NumberFormat.integer if float(v).is_integer() else NumberFormat.decimal
    if isinstance(v, str) and v.startswith("="):
        return None
    return NumberFormat.text


# Currency symbols recognised in number formats (longest first so "CN¥" wins over "¥"). "$" is the USD default.
_FORMAT_CURRENCIES = [("cn¥", "CNY"), ("₹", "INR"), ("€", "EUR"), ("£", "GBP"), ("¥", "JPY")]
_DEFAULT_DECIMALS = {NumberFormat.decimal: 2, NumberFormat.currency: 2, NumberFormat.percent: 1}


def _currency_and_decimals(number_format: str, fmt: NumberFormat | None) -> tuple[str | None, int | None]:
    """Currency code and decimal places from an Excel number format; None where the default applies."""
    if fmt not in _DEFAULT_DECIMALS:
        return None, None
    nf = number_format.lower()
    currency = None
    if fmt == NumberFormat.currency:
        currency = next((code for sym, code in _FORMAT_CURRENCIES if sym in nf), None)
    # Decimal places: zeros after the point in the positive section, ignoring quoted text and [..] blocks.
    positive = re.sub(r'"[^"]*"|\[[^\]]*\]', "", nf.split(";")[0])
    m = re.search(r"\.(0+)", positive)
    decimals = len(m.group(1)) if m else 0
    default = 0 if currency == "JPY" else _DEFAULT_DECIMALS[fmt]
    return currency, None if decimals == default else decimals


def _freeze(cell: str | None) -> tuple[bool, int]:
    """'C2' -> (True, 2): header row frozen plus two leading columns."""
    m = re.match(r"^\$?([A-Za-z]{1,3})\$?(\d+)$", cell or "")
    if not m:
        return False, 0
    return int(m.group(2)) > 1, column_index_from_string(m.group(1).upper()) - 1


def _styled(cell: CellSpec) -> bool:
    return bool(cell.fill_color or cell.bold or cell.italic or cell.font_color or cell.wrap)


def _row_height(ws, row: int) -> float | None:
    dim = ws.row_dimensions.get(row)  # .get: indexing would create the row
    if dim is None or not dim.customHeight or not dim.ht:
        return None
    return round(float(dim.ht), 1)


def _merges(ws, max_row: int, max_col: int) -> list[str]:
    """Merged ranges inside the imported area."""
    return [
        rng.coord for rng in ws.merged_cells.ranges
        if rng.max_row <= max_row and rng.max_col <= max_col
    ]


def _filters(ws, max_col: int) -> list[FilterSpec]:
    """Autofilter value selections (best effort: custom, top-10 and colour filters are skipped)."""
    if not ws.auto_filter.ref:
        return []
    try:
        first_col = range_boundaries(ws.auto_filter.ref)[0]
    except ValueError:
        return []
    out = []
    for fc in ws.auto_filter.filterColumn:
        values = [str(v) for v in (fc.filters.filter if fc.filters else [])]
        column = first_col - 1 + fc.colId
        if values and column < max_col:
            out.append(FilterSpec(column=column, values=values))
    return out


MAX_LOCKED_RANGES = 50


def _protection(ws, max_row: int, max_col: int) -> tuple[bool, list[str]]:
    """A protected sheet: fully locked when no cell is unlocked, else the locked cells as A1 rectangles.

    Best effort: more than MAX_LOCKED_RANGES rectangles (a scattered pattern) imports as unprotected.
    """
    if not ws.protection.sheet:
        return False, []
    grid = [[ws.cell(row=r, column=c).protection.locked is not False for c in range(1, max_col + 1)]
            for r in range(1, max_row + 1)]
    if all(all(row) for row in grid):
        return True, []
    # Runs of locked cells per row, then stack identical runs on consecutive rows into rectangles.
    open_boxes: dict[tuple[int, int], list[int]] = {}  # (c1, c2) -> [r1, r2]
    boxes: list[tuple[int, int, int, int]] = []
    for r, row in enumerate(grid, start=1):
        runs, c = [], 0
        while c < len(row):
            if row[c]:
                start = c
                while c < len(row) and row[c]:
                    c += 1
                runs.append((start + 1, c))
            else:
                c += 1
        for key in list(open_boxes):
            if key not in runs or open_boxes[key][1] != r - 1:
                r1, r2 = open_boxes.pop(key)
                boxes.append((key[0], r1, key[1], r2))
        for key in runs:
            if key in open_boxes:
                open_boxes[key][1] = r
            else:
                open_boxes[key] = [r, r]
        if len(boxes) + len(open_boxes) > MAX_LOCKED_RANGES:
            return False, []
    boxes += [(k[0], v[0], k[1], v[1]) for k, v in open_boxes.items()]
    boxes.sort(key=lambda b: (b[1], b[0]))
    out = []
    for c1, r1, c2, r2 in boxes:
        a, b = f"{get_column_letter(c1)}{r1}", f"{get_column_letter(c2)}{r2}"
        out.append(a if a == b else f"{a}:{b}")
    return False, out


def _row_kind(cells: list[CellSpec]) -> RowKind:
    first = next((c.value for c in cells if c.value not in (None, "")), None)
    if isinstance(first, str) and re.match(r"^(grand\s+)?total\b", first.strip(), re.I):
        return RowKind.total
    if isinstance(first, str) and re.match(r"^sub-?total\b", first.strip(), re.I):
        return RowKind.subtotal
    return RowKind.data


def _validations(ws) -> list[DataValidationSpec]:
    out: list[DataValidationSpec] = []
    for dv in ws.data_validations.dataValidation:
        rng = str(dv.sqref).split()[0]
        if dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
            options = [o.strip() for o in dv.formula1.strip('"').split(",") if o.strip()]
            out.append(DataValidationSpec(range=rng, type=ValidationType.list, options=options, message=dv.error))
        elif dv.type in ("whole", "decimal", "date"):
            def num(f):
                try:
                    return float(f) if f is not None else None
                except ValueError:
                    return None
            lo, hi = num(dv.formula1), num(dv.formula2)
            if dv.operator in ("lessThanOrEqual", "lessThan"):
                lo, hi = None, lo
            out.append(DataValidationSpec(range=rng, type=ValidationType(dv.type), min=lo, max=hi, message=dv.error))
    return out


# ---------------------------------------------------------------------------
# csv
# ---------------------------------------------------------------------------

_NUM = re.compile(r"^[-+]?[$€£]?\s?[-+]?\d{1,3}(,\d{3})*(\.\d+)?%?$|^[-+]?[$€£]?\d+(\.\d+)?%?$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _import_csv(title: str, data: bytes) -> tuple[WorkbookSpec, list[Issue]]:
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text or not text.strip():
        raise ImportError_("That file is empty.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    table = [row for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]
    issues: list[Issue] = []
    if len(table) > MAX_ROWS + 1:
        issues.append(Issue(severity="warning", message=f"Only the first {MAX_ROWS} rows were imported."))
        table = table[: MAX_ROWS + 1]

    header, body = table[0], table[1:]
    ncols = min(MAX_COLS, max(len(r) for r in table))
    columns = [ColumnSpec(header=(header[i].strip() if i < len(header) and header[i].strip() else f"Column {get_column_letter(i + 1)}"))
               for i in range(ncols)]
    rows: list[RowSpec] = []
    votes: list[dict[NumberFormat, int]] = [{} for _ in columns]
    for raw in body:
        cells = []
        for i, txt in enumerate(raw[:ncols]):
            value, fmt = _parse_csv_value(txt)
            cells.append(CellSpec(value=value))
            if value is not None:
                votes[i][fmt] = votes[i].get(fmt, 0) + 1
        rows.append(RowSpec(kind=_row_kind(cells), cells=cells))
    for i, v in enumerate(votes):
        if v:
            columns[i].format = max(v, key=v.get)
    return WorkbookSpec(title=title, sheets=[SheetSpec(name=title[:31] or "Data", columns=columns, rows=rows)]), issues


def _parse_csv_value(txt: str):
    s = txt.strip()
    if not s:
        return None, NumberFormat.text
    if s.startswith("="):
        return s, NumberFormat.text
    if _NUM.match(s):
        n = float(re.sub(r"[$€£,\s%+]", "", s))
        if s.endswith("%"):
            return n / 100, NumberFormat.percent
        if s.lstrip("-+")[:1] in "$€£":
            return n, NumberFormat.currency
        return n, NumberFormat.integer if n.is_integer() and "." not in s else NumberFormat.decimal
    if _DATE.match(s):
        return s, NumberFormat.date
    return txt, NumberFormat.text


# ---------------------------------------------------------------------------
# Tables read from a photo or PDF (see llm_engine.extract)
# ---------------------------------------------------------------------------
MAX_EXTRACTED_TABLES = 5
_NUMERIC = {NumberFormat.integer, NumberFormat.decimal, NumberFormat.currency, NumberFormat.percent}
_LOOSE_DATE = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")


def tables_to_workbook(result) -> WorkbookSpec:
    """An ExtractResult -> one sheet per table: numbers as numbers, dates as ISO strings with the date format."""
    from app.services.formula_validator import sanitize_sheet_name

    sheets: list[SheetSpec] = []
    taken: set[str] = set()
    for table in result.tables[:MAX_EXTRACTED_TABLES]:
        if not table.columns:
            continue
        name = sanitize_sheet_name(table.name or f"Table {len(sheets) + 1}", taken)
        taken.add(name.lower())
        columns = [ColumnSpec(header=c.header.strip() or f"Column {get_column_letter(i + 1)}", format=c.format,
                              currency=(c.currency or "").upper() or None if c.format == NumberFormat.currency else None)
                   for i, c in enumerate(table.columns[:MAX_COLS])]
        rows = []
        for row in table.rows[:MAX_ROWS]:
            cells = [CellSpec(value=_extracted_value(v, columns[i].format)) for i, v in enumerate(row.cells[:len(columns)])]
            rows.append(RowSpec(kind=row.kind if row.kind != RowKind.section else RowKind.data, cells=cells))
        sheets.append(SheetSpec(name=name, columns=columns, rows=rows))
    if not sheets:
        raise ImportError_("No tables were found in that file. Try a sharper photo, or crop it to the table.")
    title = (getattr(result, "title", "") or "").strip()[:100] or sheets[0].name
    return WorkbookSpec(title=title, sheets=sheets)


def _extracted_value(v, fmt: NumberFormat):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, str) and v.strip().startswith("="):
        return "'" + v.strip()  # text that looks like a formula stays text
    if fmt in _NUMERIC and isinstance(v, str):
        s = v.strip()
        negative = s.startswith("(") and s.endswith(")") or s.endswith("-")
        s = s.strip("()").rstrip("-")
        parsed, _ = _parse_csv_value(re.sub(r"[^\d.,%+\-$€£]", "", s))
        if isinstance(parsed, float):
            return -abs(parsed) if negative else parsed
        return v
    if fmt == NumberFormat.date and isinstance(v, str):
        m = _LOOSE_DATE.match(v.strip())
        if m:
            try:
                return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
            except ValueError:
                return v
    return v
