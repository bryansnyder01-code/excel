"""Charts, dropdowns, import, templates, pivot summaries, explain and PDF export."""

import io
import json
import shutil

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routes import get_engine
from app.main import app
from app.schemas import Aggregation, PivotRequest, WorkbookSpec
from app.services.importer import import_file
from app.services.pivot import PivotError, build_pivot
from app.services.workbook_builder import workbook_to_bytes
from tests.test_api import make_engine

client = TestClient(app)


def roundtrip(spec: WorkbookSpec):
    return load_workbook(io.BytesIO(workbook_to_bytes(spec)))


def sales_spec() -> WorkbookSpec:
    data = [("East", "Q1", 100), ("West", "Q1", 80), ("East", "Q2", 120), ("North", "Q2", 50), ("West", "Q2", 90)]
    return WorkbookSpec.model_validate({"title": "Sales", "sheets": [{
        "name": "Data",
        "columns": [{"header": "Region"}, {"header": "Quarter"}, {"header": "Revenue", "format": "currency"}],
        "rows": [{"cells": [{"value": r}, {"value": q}, {"value": v}]} for r, q, v in data]
               + [{"kind": "total", "cells": [{"value": "Total"}, {}, {"value": "=SUM(C2:C6)"}]}],
    }]})


def test_charts_and_dropdowns_are_written(budget_spec):
    d = budget_spec.model_dump()
    d["sheets"][2]["charts"] = [{"type": "column", "title": "Profit", "categories": "A2:A7",
                                     "series": [{"name": "Net", "range": "D2:D7"}]}]
    d["sheets"][0]["validations"] = [
        {"range": "A2:A20", "type": "list", "options": ["Product sales", "Consulting", "Other"]},
        {"range": "B2:G20", "type": "decimal", "min": 0, "message": "No negatives"},
    ]
    wb = roundtrip(WorkbookSpec.model_validate(d))
    assert len(wb["Summary"]._charts) == 1
    dvs = wb["Income"].data_validations.dataValidation
    assert {d.type for d in dvs} == {"list", "decimal"}
    assert any(d.formula1 == '"Product sales,Consulting,Other"' for d in dvs)


def test_long_dropdown_uses_hidden_sheet():
    spec = sales_spec()
    spec.sheets[0].validations = WorkbookSpec.model_validate({"title": "x", "sheets": [{
        "name": "x", "columns": [], "rows": [],
        "validations": [{"range": "A2:A50", "type": "list", "options": [f"Option, number {i}" for i in range(40)]}]}]}).sheets[0].validations
    wb = roundtrip(spec)
    assert wb["_lists"].sheet_state == "hidden"
    assert wb["Data"].data_validations.dataValidation[0].formula1.startswith("'_lists'!")


def test_import_roundtrip_keeps_formulas_formats_and_dropdowns(budget_spec):
    budget_spec.sheets[0].validations = WorkbookSpec.model_validate({"title": "x", "sheets": [{
        "name": "x", "columns": [], "rows": [],
        "validations": [{"range": "A2:A9", "type": "list", "options": ["a", "b"]}]}]}).sheets[0].validations
    budget_spec.sheets[0].rows[0].cells[0].fill_color = "#FFF2CC"
    spec, issues = import_file("budget.xlsx", workbook_to_bytes(budget_spec))
    inc = spec.sheets[0]
    assert [s.name for s in spec.sheets] == ["Income", "Expenses", "Summary"]
    assert inc.columns[1].header == "Jan" and inc.columns[1].format.value == "currency"
    assert inc.rows[0].cells[7].value == "=SUM(B2:G2)"
    assert inc.rows[3].kind.value == "total"
    assert inc.rows[0].cells[0].fill_color == "#FFF2CC"
    assert inc.rows[1].cells[0].fill_color is None  # theme banding is not frozen in
    assert inc.validations[0].options == ["a", "b"]


def test_import_csv_detects_numbers_and_formats():
    csv = "Date,Item,Amount,Share\n2026-01-05,Coffee,\"$1,200.50\",15%\n2026-01-06,Tea,300,5%\n"
    spec, _ = import_file("costs.csv", csv.encode())
    s = spec.sheets[0]
    assert [c.format.value for c in s.columns] == ["date", "text", "currency", "percent"]
    assert s.rows[0].cells[2].value == 1200.5 and s.rows[0].cells[3].value == pytest.approx(0.15)


def test_import_endpoint_rejects_other_files():
    r = client.post("/api/workbooks/import", files={"file": ("notes.docx", b"abc")})
    assert r.status_code == 422


def test_templates_list_and_build():
    r = client.get("/api/templates")
    assert r.status_code == 200 and len(r.json()) == 7
    for t in r.json():
        wb = client.get(f"/api/templates/{t['id']}").json()
        x = client.post("/api/workbooks/export", json=wb)
        assert x.status_code == 200, t["id"]
        load_workbook(io.BytesIO(x.content))


def test_pivot_by_region_and_quarter():
    spec = sales_spec()
    out, summary = build_pivot(spec, PivotRequest(workbook=spec, source_sheet="Data", row_field="Region",
                                                  column_field="Quarter", value_field="Revenue"))
    p = out.sheets[-1]
    assert [c.header for c in p.columns] == ["Region", "Q1", "Q2", "Total"]
    assert [r.cells[0].value for r in p.rows] == ["East", "North", "West", "Grand total"]
    assert p.rows[0].cells[1].value == "=SUMIFS('Data'!$C$2:$C$6,'Data'!$A$2:$A$6,$A2,'Data'!$B$2:$B$6,B$1)"
    assert p.rows[-1].cells[3].value == "=SUM('Data'!$C$2:$C$6)"
    assert p.charts and "Region" in summary


def test_pivot_count_and_errors():
    spec = sales_spec()
    out, _ = build_pivot(spec, PivotRequest(workbook=spec, source_sheet="Data", row_field="Region", aggregation=Aggregation.count))
    assert out.sheets[-1].rows[0].cells[1].value == "=COUNTIFS('Data'!$A$2:$A$6,$A2)"
    with pytest.raises(PivotError):
        build_pivot(spec, PivotRequest(workbook=spec, source_sheet="Data", row_field="Nope", value_field="Revenue"))


def test_pivot_endpoint():
    spec = sales_spec()
    r = client.post("/api/workbooks/pivot", json={"workbook": spec.model_dump(), "source_sheet": "Data",
                                                   "row_field": "Region", "value_field": "Revenue", "aggregation": "average"})
    assert r.status_code == 200
    assert r.json()["workbook"]["sheets"][-1]["rows"][0]["cells"][1]["value"].startswith("=IFERROR(AVERAGEIFS(")


def test_explain_endpoint_sends_cell_context():
    seen: list[dict] = []
    reply = {"explanation": "It divides profit by income, but income is empty.", "fixed_value": "=IFERROR(D2/B2,0)"}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply], seen)
    try:
        spec = sales_spec()
        r = client.post("/api/workbooks/explain", json={"workbook": spec.model_dump(), "sheet": "Data",
                                                         "cell": "C7", "computed": "440"})
        assert r.status_code == 200, r.text
        assert r.json()["fixed_value"] == "=IFERROR(D2/B2,0)"
        prompt = seen[0]["body"]["messages"][0]["content"]
        assert "Cell contents: =SUM(C2:C6)" in prompt and "Currently shows: 440" in prompt
        assert seen[0]["body"]["output_config"]["effort"] == "low"
    finally:
        app.dependency_overrides.clear()


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_pdf_export(budget_spec):
    r = client.post("/api/workbooks/export?format=pdf", json=budget_spec.model_dump())
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    assert r.headers["content-disposition"].endswith('.pdf"')
