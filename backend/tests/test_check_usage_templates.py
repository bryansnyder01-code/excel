"""Layout and chart options, scenarios, "check my workbook", AI usage limits and user templates."""

import io

import anthropic
import httpx2
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import my_templates
from app.api.routes import get_engine
from app.config import Settings, get_settings
from app.main import app
from app.schemas import CheckRequest, WorkbookSpec
from app.services.checker import deterministic_findings, finding_id, formula_shape, move_formula
from app.services.formula_validator import validate_workbook
from app.services.importer import import_file
from app.services.llm_engine import WorkbookEngine
from app.services.workbook_builder import workbook_to_bytes
from tests.test_api import make_engine, tiny_workbook
from tests.test_more import auth, signup

client = TestClient(app)


def layout_spec() -> WorkbookSpec:
    return WorkbookSpec.model_validate({
        "title": "Regions",
        "scenarios": [{"name": "Base", "values": {"Data!B2": 5, "'Data'!$c$2": 0.1}},
                      {"name": "High", "values": {"'Data'!B2": 9}}],
        "sheets": [{
            "name": "Data", "header_height": 30, "merges": ["A1:B1", "A6:C6"],
            "filters": [{"column": 0, "values": ["East", "North"]}],
            "columns": [{"header": "Region"}, {"header": "Q1", "format": "currency"}, {"header": "Q2", "format": "currency"}],
            "rows": [
                {"height": 40, "cells": [{"value": "East", "wrap": True, "align": "center"}, {"value": 1}, {"value": 2}]},
                {"cells": [{"value": "West"}, {"value": 3}, {"value": 4}]},
                {"cells": [{"value": '="No"&"rth"'}, {"value": 3}, {"value": 4}]},
                {"cells": [{"value": "north"}, {"value": 5}, {"value": 6}]},
                {"cells": [{"value": "Notes span the row", "wrap": True}]},
                {"kind": "total", "cells": [{"value": "Total"}, {"value": "=SUM(B2:B5)"}, {"value": "=SUM(C2:C5)"}]},
            ],
            "charts": [
                {"type": "column", "title": "Combo", "categories": "A2:A5", "stacked": True, "x_title": "Region",
                 "y_title": "Sales", "series": [{"name": "Q1", "range": "B2:B5", "color": "#C00000"},
                                                {"name": "Q2", "range": "C2:C5", "kind": "line", "color": "00B050"}]},
                {"type": "scatter", "title": "XY", "categories": "B2:B5", "series": [{"name": "Q2", "range": "C2:C5"}]},
                {"type": "area", "title": "Area", "categories": "A2:A5", "stacked": True,
                 "series": [{"name": "Q1", "range": "B2:B5"}, {"name": "Q2", "range": "C2:C5"}]},
                {"type": "doughnut", "title": "Share", "categories": "A2:A5", "series": [{"name": "Q2", "range": "C2:C5"}]},
            ],
        }],
    })


# ---------------------------------------------------------------------------
# Builder and importer
# ---------------------------------------------------------------------------


def test_builder_layout_merges_wrap_heights_and_filters():
    wb = load_workbook(io.BytesIO(workbook_to_bytes(layout_spec())))
    ws = wb["Data"]
    assert {str(r) for r in ws.merged_cells.ranges} == {"A1:B1", "A6:C6"}
    assert ws.row_dimensions[1].height == 30 and ws.row_dimensions[2].height == 40
    assert ws["A2"].alignment.wrap_text and ws["A2"].alignment.horizontal == "center"
    assert not ws["A3"].alignment.wrap_text

    (fc,) = ws.auto_filter.filterColumn
    assert fc.colId == 0 and fc.filters.filter == ["East", "North"]
    hidden = {r for r in range(2, 8) if ws.row_dimensions[r].hidden}
    assert hidden == {3, 6}  # West and the notes row; formula row and "north" (case-insensitive) stay; total untouched


def test_builder_new_charts_and_scenarios_sheet():
    data = workbook_to_bytes(layout_spec())
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Data"]
    assert [type(c).__name__ for c in ws._charts] == ["BarChart", "ScatterChart", "AreaChart", "DoughnutChart"]

    import zipfile
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/charts/chart1.xml").decode()
    assert "lineChart" in xml and "barChart" in xml  # combo chart
    assert 'grouping val="stacked"' in xml and 'overlap val="100"' in xml
    assert "C00000" in xml and "00B050" in xml
    assert "Region" in xml and "Sales" in xml  # axis titles
    area = zipfile.ZipFile(io.BytesIO(data)).read("xl/charts/chart3.xml").decode()
    assert 'grouping val="stacked"' in area

    assert wb.sheetnames == ["Data", "Scenarios"]
    rows = [[c.value for c in r] for r in wb["Scenarios"].iter_rows()]
    assert rows == [["Input cell", "Base", "High"], ["'Data'!B2", 5, 9], ["'Data'!C2", 0.1, None]]


def test_import_roundtrip_keeps_merges_wrap_heights_and_filters():
    spec, issues = import_file("regions.xlsx", workbook_to_bytes(layout_spec()))
    sheet = spec.sheets[0]
    assert sorted(sheet.merges) == ["A1:B1", "A6:C6"]
    assert not any("Merged" in i.message for i in issues)
    assert sheet.header_height == 30
    assert sheet.rows[0].height == 40 and sheet.rows[1].height is None
    assert sheet.rows[0].cells[0].wrap and not sheet.rows[1].cells[0].wrap
    assert [(f.column, f.values) for f in sheet.filters] == [(0, ["East", "North"])]


def test_validator_warns_about_bad_merges_filters_and_chart_options():
    spec = layout_spec()
    sheet = spec.sheets[0]
    sheet.merges += ["B1:C2", "nonsense"]
    sheet.filters.append(sheet.filters[0].model_copy(update={"column": 9}))
    sheet.charts[1].stacked = True
    sheet.charts[3].series[0].color = "red"
    messages = [i.message for i in validate_workbook(spec)]
    assert any("overlaps another merge" in m for m in messages)
    assert any('"nonsense" is not a valid' in m for m in messages)
    assert any("filter points at column 10" in m for m in messages)
    assert any("can be stacked" in m for m in messages)
    assert any('colour "red"' in m for m in messages)
    workbook_to_bytes(spec)  # bad entries are skipped, not fatal


# ---------------------------------------------------------------------------
# Checker
# ---------------------------------------------------------------------------


def check_spec() -> WorkbookSpec:
    return WorkbookSpec.model_validate({"title": "Shop", "sheets": [
        {"name": "Sales", "columns": [{"header": "Item"}, {"header": "Qty"}, {"header": "Price"}, {"header": "Revenue"}],
         "rows": [
             {"cells": [{"value": "A"}, {"value": 1}, {"value": 10}, {"value": "=B2*C2"}]},
             {"cells": [{"value": "B"}, {}, {"value": 20}, {"value": "=B3*C3"}]},
             {"cells": [{"value": "C"}, {"value": 3}, {"value": 0}, {"value": "=B4*C4"}]},
             {"cells": [{"value": "D"}, {"value": 4}, {"value": 40}, {"value": "=B5*C4"}]},
             {"cells": [{"value": "E"}, {"value": 5}, {"value": 50}, {"value": 300}]},
             {"kind": "total", "cells": [{"value": "Total"}, {"value": "=SUM(B2:B6)"}, {}, {"value": 1234}]},
         ]},
        {"name": "Summary", "columns": [{"header": "What"}, {"header": "Value"}],
         "rows": [{"cells": [{"value": "Lost"}, {"value": "=Missing!A1"}]}]},
    ]})


CHECK_VALUES = {"Sales": [["Item", "Qty", "Price", "Revenue"], ["A", "1", "10", "10"], ["B", "", "20", "0"],
                          ["C", "3", "0", "#DIV/0!"]]}


def by_cell(findings):
    return {(f.sheet, f.cell): f for f in findings}


def test_formula_shape_and_move():
    assert formula_shape("=B2*C2", 2, 4) == formula_shape("=B9*C9", 9, 4)
    assert formula_shape("=B5*C4", 5, 4) != formula_shape("=B2*C2", 2, 4)
    assert formula_shape("=B2*$C$1", 2, 4) == formula_shape("=B7*$C$1", 7, 4)
    assert move_formula("=B2*$C$1+'Q1 Data'!A2&\"B2\"", 2, 6) == "=B6*$C$1+'Q1 Data'!A6&\"B2\""
    assert move_formula("=SUM(B$2:B2)", 2, 4) == "=SUM(B$2:B4)"


def test_deterministic_rules():
    findings = deterministic_findings(check_spec(), CHECK_VALUES)
    cells = {}
    for f in findings:
        cells.setdefault((f.sheet, f.cell), []).append(f)

    (err,) = cells[("Sales", "D4")]
    assert err.severity == "error" and err.kind == "excel_error" and "#DIV/0!" in err.message and err.fix is None

    total = next(f for f in cells[("Sales", "D7")] if f.kind == "hardcoded_total")
    assert total.severity == "warning" and total.fix == "=SUM(D2:D6)"

    odd = next(f for f in cells[("Sales", "D5")] if f.kind == "inconsistent_formula")
    assert odd.fix == "=B5*C5"

    typed = next(f for f in cells[("Sales", "D6")] if f.kind == "number_in_formula_column")
    assert typed.fix == "=B6*C6"

    empty = next(f for f in cells[("Sales", "D3")] if f.kind == "references_empty")
    assert empty.severity == "info" and empty.fix is None and "B3" in empty.message

    missing = next(f for f in cells[("Summary", "B2")] if f.kind == "validation")
    assert missing.severity == "error" and "Missing" in missing.message

    assert total.id == finding_id("hardcoded_total", "Sales", "D7")
    assert [f.id for f in deterministic_findings(check_spec(), CHECK_VALUES)] == [f.id for f in findings]


def test_check_endpoint_without_ai_does_not_call_claude():
    seen: list[dict] = []
    app.dependency_overrides[get_engine] = lambda: make_engine([], seen)
    try:
        r = client.post("/api/workbooks/check", json={"workbook": check_spec().model_dump(), "values": CHECK_VALUES,
                                                      "use_ai": False})
    finally:
        app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    body = r.json()
    assert seen == []
    assert body["summary"] == "Found 2 errors, 3 warnings and 1 note."
    sev = [f["severity"] for f in body["findings"]]
    assert sev == sorted(sev, key=["error", "warning", "info"].index)
    assert len({(f["sheet"], f["cell"]) for f in body["findings"]}) == len(body["findings"])
    assert set(body["findings"][0]) == {"id", "severity", "sheet", "cell", "kind", "message", "fix"}


def test_check_merges_ai_findings():
    seen: list[dict] = []
    reply = {"summary": "Two problems.", "findings": [
        {"severity": "error", "sheet": "Sales", "cell": "$d$7", "kind": "double count", "message": "Counts twice.",
         "fix": "=SUM(D2:D6)"},
        {"severity": "warning", "sheet": "sales", "cell": "'Sales'!B7", "kind": "range_misses_rows",
         "message": "Misses a row.", "fix": "=SUM(B2:B6)"},
        {"severity": "error", "sheet": "Nope", "cell": "A1", "kind": "x", "message": "Dropped.", "fix": None},
    ]}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply], seen)
    try:
        r = client.post("/api/workbooks/check", json={"workbook": check_spec().model_dump(), "values": CHECK_VALUES})
    finally:
        app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    found = {(f["sheet"], f["cell"]): f for f in r.json()["findings"]}
    assert found[("Sales", "D7")]["kind"] == "double_count" and found[("Sales", "D7")]["severity"] == "error"
    assert found[("Sales", "B7")]["fix"] == "=SUM(B2:B6)"
    assert not any(s == "Nope" for s, _ in found)
    assert r.json()["summary"] == "Found 3 errors, 3 warnings and 1 note."

    sent = seen[0]["body"]
    assert sent["output_config"]["effort"] == "low"
    assert sent["output_config"]["format"]["type"] == "json_schema"
    prompt = sent["messages"][0]["content"]
    assert "<already_found>" in prompt and "'Sales'!D4" in prompt and "  4 | C | 3 | 0 | #DIV/0!" in prompt


def test_check_falls_back_when_ai_fails():
    transport = httpx2.MockTransport(lambda request: httpx2.Response(
        529, json={"type": "error", "error": {"type": "overloaded_error", "message": "busy"}}))
    busy = anthropic.AsyncAnthropic(api_key="test", max_retries=0, http_client=httpx2.AsyncClient(transport=transport))
    app.dependency_overrides[get_engine] = lambda: WorkbookEngine(client=busy, settings=Settings())
    try:
        r = client.post("/api/workbooks/check", json={"workbook": check_spec().model_dump(), "values": CHECK_VALUES})
    finally:
        app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"].startswith("Found 2 errors") and "AI review was unavailable" in body["summary"]
    assert len(body["findings"]) == 6


# ---------------------------------------------------------------------------
# AI usage limits
# ---------------------------------------------------------------------------


@pytest.fixture
def tiny_limits(monkeypatch):
    s = get_settings()
    monkeypatch.setattr(s, "ai_limit_anonymous", 2)
    monkeypatch.setattr(s, "ai_limit_signed_in", 3)
    monkeypatch.setattr(s, "ai_limit_enabled", True)


def ask_body() -> dict:
    return {"workbook": tiny_workbook("=B2*0.15"), "question": "What is the tax?"}


def test_usage_limit_returns_429(tiny_limits):
    reply = {"answer": "15.", "references": []}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply] * 10, [])
    try:
        assert client.get("/api/workbooks/usage").json() == {"used": 0, "limit": 2, "signed_in": False}
        assert client.post("/api/workbooks/ask", json=ask_body()).status_code == 200
        # A check without AI is free.
        assert client.post("/api/workbooks/check", json={"workbook": tiny_workbook("=1"), "use_ai": False}).status_code == 200
        assert client.post("/api/workbooks/ask", json=ask_body()).status_code == 200
        r = client.post("/api/workbooks/ask", json=ask_body())
        assert r.status_code == 429
        assert r.json() == {"detail": "You've used today's AI requests (2). Sign in for more, or try again tomorrow."}

        stream = client.post("/api/workbooks/generate/stream", json={"prompt": "Budget"})
        assert stream.status_code == 429 and not stream.headers["content-type"].startswith("text/event-stream")
        assert client.post("/api/workbooks/check", json={"workbook": tiny_workbook("=1")}).status_code == 429
        assert client.get("/api/workbooks/usage").json() == {"used": 2, "limit": 2, "signed_in": False}

        # Signed in: its own, larger allowance.
        h = auth(signup("limits@example.com")["token"])
        assert client.get("/api/workbooks/usage", headers=h).json() == {"used": 0, "limit": 3, "signed_in": True}
        for _ in range(3):
            assert client.post("/api/workbooks/ask", json=ask_body(), headers=h).status_code == 200
        r = client.post("/api/workbooks/ask", json=ask_body(), headers=h)
        assert r.status_code == 429 and r.json()["detail"] == "You've used today's AI requests (3). Try again tomorrow."

        # A bad token counts against the IP.
        bad = client.post("/api/workbooks/ask", json=ask_body(), headers=auth("nope"))
        assert bad.status_code == 429 and "Sign in" in bad.json()["detail"]
    finally:
        app.dependency_overrides.clear()


def test_usage_limit_can_be_disabled_and_honours_proxy_header(tiny_limits, monkeypatch):
    s = get_settings()
    monkeypatch.setattr(s, "trust_proxy", True)
    reply = {"answer": "15.", "references": []}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply] * 10, [])
    try:
        for ip in ("1.1.1.1", "1.1.1.1"):
            assert client.post("/api/workbooks/ask", json=ask_body(), headers={"X-Forwarded-For": f"{ip}, 10.0.0.1"}).status_code == 200
        assert client.post("/api/workbooks/ask", json=ask_body(), headers={"X-Forwarded-For": "1.1.1.1"}).status_code == 429
        assert client.post("/api/workbooks/ask", json=ask_body(), headers={"X-Forwarded-For": "2.2.2.2"}).status_code == 200
        monkeypatch.setattr(s, "ai_limit_enabled", False)
        assert client.post("/api/workbooks/ask", json=ask_body(), headers={"X-Forwarded-For": "1.1.1.1"}).status_code == 200
    finally:
        app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# User templates
# ---------------------------------------------------------------------------


def test_user_templates_crud_auth_and_cap(monkeypatch):
    assert client.get("/api/templates/mine").status_code == 401
    assert client.post("/api/templates/mine", json={"title": "x", "workbook": tiny_workbook("=1")}).status_code == 401

    a = auth(signup("tpl-a@example.com")["token"])
    b = auth(signup("tpl-b@example.com")["token"])
    assert client.get("/api/templates/mine", headers=a).json() == []

    r = client.post("/api/templates/mine", headers=a,
                    json={"title": "Tax sheet", "description": "Sales with tax", "workbook": layout_spec().model_dump()})
    assert r.status_code == 200, r.text
    tid = r.json()["id"]
    (listed,) = client.get("/api/templates/mine", headers=a).json()
    assert {k: listed[k] for k in ("id", "title", "description", "sheets")} == {
        "id": tid, "title": "Tax sheet", "description": "Sales with tax", "sheets": 1}
    assert isinstance(listed["created"], int)

    got = client.get(f"/api/templates/mine/{tid}", headers=a)
    assert got.status_code == 200
    assert WorkbookSpec.model_validate(got.json()) == layout_spec()

    assert client.get(f"/api/templates/mine/{tid}", headers=b).status_code == 404
    assert client.delete(f"/api/templates/mine/{tid}", headers=b).status_code == 404
    assert client.get("/api/templates/mine", headers=b).json() == []

    bad = client.post("/api/templates/mine", headers=a, json={"title": "x", "workbook": {"title": "no sheets"}})
    assert bad.status_code == 422

    assert client.delete(f"/api/templates/mine/{tid}", headers=a).status_code == 204
    assert client.get(f"/api/templates/mine/{tid}", headers=a).status_code == 404

    monkeypatch.setattr(my_templates, "MAX_TEMPLATES", 2)
    for i in range(2):
        assert client.post("/api/templates/mine", headers=a, json={"title": f"T{i}", "workbook": tiny_workbook("=1")}).status_code == 200
    full = client.post("/api/templates/mine", headers=a, json={"title": "T3", "workbook": tiny_workbook("=1")})
    assert full.status_code == 422 and full.json() == {"detail": "You can keep up to 2 templates."}

    big ={"title": "Big", "workbook": tiny_workbook("x" * (5 * 1024 * 1024))}
    assert client.post("/api/templates/mine", headers=b, json=big).status_code == 413

    # The built-in template routes still work alongside /mine.
    assert client.get("/api/templates").status_code == 200
    assert client.get("/api/templates/does-not-exist").status_code == 404


def test_edit_keeps_scenarios_the_model_cannot_return():
    original = tiny_workbook("=B2*0.15") | {"scenarios": [{"name": "Base", "values": {"Sales!B2": 100}}]}
    reply = {"summary": "Renamed.", "workbook": tiny_workbook("=B2*0.2") | {"scenarios": [{"name": "Base", "values": {}}]}}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply], [])
    try:
        r = client.post("/api/workbooks/edit", json={"prompt": "Use 20% tax", "workbook": original})
    finally:
        app.dependency_overrides.clear()
    assert r.status_code == 200, r.text
    assert r.json()["workbook"]["scenarios"] == [{"name": "Base", "values": {"Sales!B2": 100.0}}]
