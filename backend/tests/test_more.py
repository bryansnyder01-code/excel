"""Number formats, freeze panes, new conditional formats, streaming progress, ask, accounts and sharing."""

import io
import json

import anthropic
import httpx2
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.routes import get_engine
from app.config import Settings
from app.main import app
from app.schemas import NumberFormat, WorkbookSpec
from app.services.importer import import_file
from app.services.llm_engine import WorkbookEngine
from app.services.workbook_builder import number_format, workbook_to_bytes
from tests.test_api import make_engine, tiny_workbook

client = TestClient(app)


def roundtrip(spec: WorkbookSpec):
    return load_workbook(io.BytesIO(workbook_to_bytes(spec)))


def wide_spec() -> WorkbookSpec:
    return WorkbookSpec.model_validate({"title": "Costs", "sheets": [{
        "name": "Costs", "freeze_columns": 2,
        "columns": [{"header": "Item"}, {"header": "Status"},
                    {"header": "Cost", "format": "currency", "currency": "EUR"},
                    {"header": "Rate", "format": "percent", "decimals": 2},
                    {"header": "Qty", "format": "decimal", "decimals": 3}],
        "rows": [{"cells": [{"value": "Rent"}, {"value": "Overdue"}, {"value": 1200}, {"value": 0.1234}, {"value": 1.5}]},
                 {"cells": [{"value": "Power"}, {"value": "Paid"},
                            {"value": 300, "currency": "JPY"}, {"value": 0.5, "decimals": 0}, {"value": 2}]}],
        "conditional_formats": [
            {"range": "C2:C3", "operator": "color_scale", "reverse": True},
            {"range": "E2:E3", "operator": "data_bar"},
            {"range": "B2:B3", "operator": "contains", "text": 'Over"due', "style": "custom",
             "fill_color": "#FFC7CE", "font_color": "#9C0006"},
            {"range": "C2:C3", "operator": "greater_than", "value1": 1000},
        ],
    }]})


# ---------------------------------------------------------------------------
# Builder
# ---------------------------------------------------------------------------


def test_number_formats_honour_currency_and_decimals():
    assert number_format(NumberFormat.currency) == '"$"#,##0.00;[Red]-"$"#,##0.00'
    assert number_format(NumberFormat.currency, None, "EUR") == '"€"#,##0.00;[Red]-"€"#,##0.00'
    assert number_format(NumberFormat.currency, None, "JPY") == '"¥"#,##0;[Red]-"¥"#,##0'
    assert number_format(NumberFormat.currency, 0, "XYZ") == '"$"#,##0;[Red]-"$"#,##0'
    assert number_format(NumberFormat.decimal, 3) == "#,##0.000"
    assert number_format(NumberFormat.percent) == "0.0%"
    assert number_format(NumberFormat.percent, 2) == "0.00%"


def test_builder_freeze_formats_and_conditional_rules():
    ws = roundtrip(wide_spec())["Costs"]
    assert ws.freeze_panes == "C2"
    assert ws["C2"].number_format == '"€"#,##0.00;[Red]-"€"#,##0.00'
    assert ws["C3"].number_format == '"¥"#,##0;[Red]-"¥"#,##0'  # cell currency overrides the column
    assert ws["D2"].number_format == "0.00%" and ws["D3"].number_format == "0%"
    assert ws["E2"].number_format == "#,##0.000"

    rules = {r.type: r for cf in ws.conditional_formatting for r in cf.rules}
    assert set(rules) == {"colorScale", "dataBar", "expression", "cellIs"}
    colors = [c.rgb for c in rules["colorScale"].colorScale.color]
    assert colors[0].endswith("63BE7B") and colors[2].endswith("F8696B")  # reversed: low is green
    assert rules["dataBar"].dataBar.color.rgb.endswith("638EC6")
    assert rules["expression"].formula == ['ISNUMBER(SEARCH("Over""due",B2))']
    assert rules["expression"].dxf.fill.bgColor.rgb.endswith("FFC7CE")


def test_freeze_variants():
    spec = wide_spec()
    sheet = spec.sheets[0]
    sheet.freeze_columns, sheet.freeze_header = 0, False
    assert roundtrip(spec)["Costs"].freeze_panes is None
    sheet.freeze_columns = 1
    assert roundtrip(spec)["Costs"].freeze_panes == "B1"


def test_import_reads_freeze_currency_and_decimals():
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Cost", "Rate"])
    for name, cost in [("A", 10.5), ("B", 20)]:
        ws.append([name, cost, 0.25])
    for r in (2, 3):
        ws.cell(row=r, column=2).number_format = '"£"#,##0.000'
        ws.cell(row=r, column=3).number_format = "0.0%"
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    spec, _ = import_file("costs.xlsx", buf.getvalue())
    sheet = spec.sheets[0]
    assert sheet.freeze_header and sheet.freeze_columns == 1
    cost = sheet.columns[1]
    assert (cost.format, cost.currency, cost.decimals) == (NumberFormat.currency, "GBP", 3)
    assert sheet.columns[2].format == NumberFormat.percent and sheet.columns[2].decimals is None
    assert sheet.rows[0].cells[1].currency is None  # matches the column, so no per-cell override


# ---------------------------------------------------------------------------
# Streaming progress
# ---------------------------------------------------------------------------


def sse_with_thinking(payload: dict, chunk: int = 7) -> httpx2.Response:
    """Like test_api.sse_response, but with a thinking block and the JSON split into small deltas."""
    text = json.dumps(payload)
    events = [
        ("message_start", {"type": "message_start", "message": {
            "id": "msg_1", "type": "message", "role": "assistant", "model": "claude-sonnet-5-5",
            "content": [], "stop_reason": None, "stop_sequence": None,
            "usage": {"input_tokens": 10, "output_tokens": 0}}}),
        ("content_block_start", {"type": "content_block_start", "index": 0,
                                 "content_block": {"type": "thinking", "thinking": "", "signature": ""}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0,
                                 "delta": {"type": "thinking_delta", "thinking": "Two sheets."}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0,
                                 "delta": {"type": "signature_delta", "signature": "sig"}}),
        ("content_block_stop", {"type": "content_block_stop", "index": 0}),
        ("content_block_start", {"type": "content_block_start", "index": 1,
                                 "content_block": {"type": "text", "text": ""}}),
        *[("content_block_delta", {"type": "content_block_delta", "index": 1,
                                   "delta": {"type": "text_delta", "text": text[i:i + chunk]}})
          for i in range(0, len(text), chunk)],
        ("content_block_stop", {"type": "content_block_stop", "index": 1}),
        ("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None},
                           "usage": {"output_tokens": 50}}),
        ("message_stop", {"type": "message_stop"}),
    ]
    body = "".join(f"event: {e}\ndata: {json.dumps(d)}\n\n" for e, d in events)
    return httpx2.Response(200, headers={"content-type": "text/event-stream"}, content=body.encode())


def thinking_engine(replies: list[dict]) -> WorkbookEngine:
    queue = list(replies)
    transport = httpx2.MockTransport(lambda request: sse_with_thinking(queue.pop(0)))
    client_ = anthropic.AsyncAnthropic(api_key="test", http_client=httpx2.AsyncClient(transport=transport))
    return WorkbookEngine(client=client_, settings=Settings(repair_attempts=1))


def read_events(r) -> list[dict]:
    assert r.headers["content-type"].startswith("text/event-stream")
    chunks = [c for c in r.text.split("\n\n") if c.strip()]
    assert all(c.startswith("data: ") for c in chunks)
    return [json.loads(c[len("data: "):]) for c in chunks]


def test_generate_stream_reports_progress_then_result():
    wb = tiny_workbook("=Rates!B2*B2")
    wb["sheets"][0]["charts"] = [{"type": "column", "title": "Amounts", "categories": "A2:A2",
                                  "series": [{"name": "Amount", "range": "B2:B2"}]}]
    wb["sheets"].append({"name": 'Q1 "Draft"', "columns": [{"header": "X"}], "rows": []})
    good = {"summary": "Built it.", "workbook": tiny_workbook("=B2*0.15")}
    app.dependency_overrides[get_engine] = lambda: thinking_engine([{"summary": "First try.", "workbook": wb}, good])
    try:
        r = client.post("/api/workbooks/generate/stream", json={"prompt": "Sales sheet with tax"})
        assert r.status_code == 200
        events = read_events(r)
    finally:
        app.dependency_overrides.clear()

    progress = [e["text"] for e in events if e["type"] == "progress"]
    assert progress == ["Reading your request", "Planning the layout", 'Writing sheet "Sales"', "Adding charts",
                        'Writing sheet "Q1 "Draft""', "Checking formulas", "Fixing 1 formula problem"]
    assert len(progress) == len(set(progress))
    assert [e["type"] for e in events][-1] == "result" and sum(e["type"] == "result" for e in events) == 1
    final = events[-1]
    assert final["summary"] == "Built it."
    assert final["workbook"]["sheets"][0]["rows"][0]["cells"][2]["value"] == "=B2*0.15"
    assert isinstance(final["issues"], list)


def test_edit_stream_errors_are_events():
    app.dependency_overrides[get_engine] = lambda: make_engine([], [])
    try:
        long = client.post("/api/workbooks/edit/stream", json={"prompt": "x" * 5000, "workbook": tiny_workbook("=1")})
        assert read_events(long) == [{"type": "error", "message": "Please keep requests under 4000 characters."}]

        def handler(request):
            return httpx2.Response(429, json={"type": "error", "error": {"type": "rate_limit_error", "message": "slow"}})

        transport = httpx2.MockTransport(handler)
        busy = anthropic.AsyncAnthropic(api_key="test", max_retries=0, http_client=httpx2.AsyncClient(transport=transport))
        app.dependency_overrides[get_engine] = lambda: WorkbookEngine(client=busy, settings=Settings())
        r = client.post("/api/workbooks/edit/stream", json={"prompt": "Add a row", "workbook": tiny_workbook("=1")})
        events = read_events(r)
        assert events[0] == {"type": "progress", "text": "Reading your request"}
        assert events[-1]["type"] == "error" and "busy" in events[-1]["message"]
    finally:
        app.dependency_overrides.clear()


# ---------------------------------------------------------------------------
# Ask
# ---------------------------------------------------------------------------


def test_ask_endpoint_builds_table_prompt():
    seen: list[dict] = []
    reply = {"answer": "Widget is the only item: 100 with 15 tax.",
             "references": [{"sheet": "Sales", "range": f"B{i}"} for i in range(2, 9)]}
    app.dependency_overrides[get_engine] = lambda: make_engine([reply], seen)
    try:
        values = {"Sales": [["Item", "Amount", "Tax"], ["Widget", "$100.00", "$15.00"]]}
        r = client.post("/api/workbooks/ask", json={"workbook": tiny_workbook("=B2*0.15"),
                                                    "question": "What is the tax on Widget?", "values": values})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["answer"].startswith("Widget")
        assert len(body["references"]) == 5 and body["references"][0] == {"sheet": "Sales", "range": "B2"}

        sent = seen[0]["body"]
        prompt = sent["messages"][0]["content"]
        assert "    | A | B | C" in prompt and "  2 | Widget | $100.00 | $15.00" in prompt
        assert "'Sales'!C2 =B2*0.15" in prompt and "What is the tax on Widget?" in prompt
        assert sent["output_config"]["effort"] == "low" and sent["max_tokens"] == 8000
        assert "Never change the workbook" in sent["system"][0]["text"]

        empty = client.post("/api/workbooks/ask", json={"workbook": tiny_workbook("=1"), "question": "  "})
        assert empty.status_code == 422
    finally:
        app.dependency_overrides.clear()


def test_ask_truncates_long_sheets():
    from app.services.llm_engine import _ask_tables

    grid = [["N"]] + [[str(i)] for i in range(500)]
    text = _ask_tables(WorkbookSpec.model_validate(tiny_workbook("=1")), {"Sales": grid})
    assert "showing the first 300 of 501 rows" in text
    assert "300 | 298" in text and "301 |" not in text


# ---------------------------------------------------------------------------
# Accounts, sync and sharing
# ---------------------------------------------------------------------------


def signup(email: str, password: str = "correct horse") -> dict:
    r = client.post("/api/auth/signup", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return r.json()


def auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def saved_doc(doc_id: str, title: str = "Budget") -> dict:
    return {"id": doc_id, "title": title, "updated": 1700000000000,
            "workbook": tiny_workbook("=B2*0.15"), "messages": [{"role": "user", "text": "secret"}], "checkpoints": []}


def test_signup_validation_and_login():
    assert client.post("/api/auth/signup", json={"email": "nope", "password": "longenough"}).json() == {
        "detail": "Enter a valid email address."}
    r = client.post("/api/auth/signup", json={"email": "a@b.co", "password": "short"})
    assert r.status_code == 422 and r.json()["detail"] == "Use at least 8 characters for your password."

    alice = signup("  Alice@Example.com ")
    assert alice["user"]["email"] == "alice@example.com" and alice["token"]
    dup = client.post("/api/auth/signup", json={"email": "ALICE@example.com", "password": "whatever123"})
    assert dup.status_code == 409 and dup.json()["detail"] == "An account with that email already exists."

    assert client.get("/api/auth/me").status_code == 401
    me = client.get("/api/auth/me", headers=auth(alice["token"]))
    assert me.json() == alice["user"]

    wrong = client.post("/api/auth/login", json={"email": "alice@example.com", "password": "nope nope"})
    assert wrong.status_code == 401 and wrong.json()["detail"] == "Wrong email or password."
    ok = client.post("/api/auth/login", json={"email": "ALICE@example.com", "password": "correct horse"})
    assert ok.status_code == 200 and ok.json()["user"] == alice["user"]

    # Logout ends only that session.
    assert client.post("/api/auth/logout", headers=auth(ok.json()["token"])).status_code == 204
    gone = client.get("/api/auth/me", headers=auth(ok.json()["token"]))
    assert gone.status_code == 401 and gone.json()["detail"] == "Please sign in again."
    assert client.get("/api/auth/me", headers=auth(alice["token"])).status_code == 200


def test_login_throttle():
    signup("bob@example.com")
    for _ in range(5):
        assert client.post("/api/auth/login", json={"email": "bob@example.com", "password": "bad pass"}).status_code == 401
    r = client.post("/api/auth/login", json={"email": "bob@example.com", "password": "correct horse"})
    assert r.status_code == 429 and r.json()["detail"] == "Too many attempts. Try again in a few minutes."


def test_docs_sync_and_share_links():
    alice, eve = signup("alice@example.com"), signup("eve@example.com")
    a, e = auth(alice["token"]), auth(eve["token"])

    assert client.get("/api/docs").status_code == 401
    put = client.put("/api/docs/doc1", headers=a, json={"title": "Budget", "updated": 1700000000000,
                                                        "data": saved_doc("doc1")})
    assert put.status_code == 200 and put.json() == {"id": "doc1", "updated": 1700000000000}
    client.put("/api/docs/doc2", headers=a, json={"title": "Newer", "updated": 1700000005000, "data": {"x": 1}})

    listed = client.get("/api/docs", headers=a).json()
    assert [d["id"] for d in listed] == ["doc2", "doc1"]
    assert listed[1] == {"id": "doc1", "title": "Budget", "updated": 1700000000000, "sheets": 1, "shared": False,
                         "role": "owner", "owner_email": "alice@example.com"}
    assert listed[0]["sheets"] == 0

    got = client.get("/api/docs/doc1", headers=a).json()
    assert got["data"]["messages"][0]["text"] == "secret" and got["title"] == "Budget"

    # Someone else can't read, overwrite or share it.
    assert client.get("/api/docs/doc1", headers=e).json() == {"detail": "Workbook not found."}
    assert client.put("/api/docs/doc1", headers=e, json={"title": "x", "updated": 1, "data": {}}).status_code == 404
    assert client.post("/api/docs/doc1/share", headers=e).status_code == 404
    assert client.get("/api/docs", headers=e).json() == []

    # Last write wins.
    client.put("/api/docs/doc1", headers=a, json={"title": "Budget v2", "updated": 1700000009000, "data": saved_doc("doc1")})
    assert client.get("/api/docs/doc1", headers=a).json()["title"] == "Budget v2"

    # Share: same token twice, public fetch shows only the workbook.
    token = client.post("/api/docs/doc1/share", headers=a).json()["token"]
    assert client.post("/api/docs/doc1/share", headers=a).json()["token"] == token
    assert client.get("/api/docs", headers=a).json()[0]["shared"] is True
    public = client.get(f"/api/shared/{token}")
    assert public.status_code == 200
    assert public.json() == {"title": "Budget v2", "workbook": tiny_workbook("=B2*0.15")}
    assert "secret" not in public.text

    assert client.delete("/api/docs/doc1/share", headers=a).status_code == 204
    off = client.get(f"/api/shared/{token}")
    assert off.status_code == 404 and off.json()["detail"] == "This share link doesn't exist or was turned off."

    # Deleting a doc removes it and its share link.
    token2 = client.post("/api/docs/doc1/share", headers=a).json()["token"]
    assert client.delete("/api/docs/doc1", headers=a).status_code == 204
    assert client.get("/api/docs/doc1", headers=a).status_code == 404
    assert client.get(f"/api/shared/{token2}").status_code == 404


def test_doc_size_limit():
    a = auth(signup("big@example.com")["token"])
    blob = "x" * (5 * 1024 * 1024)
    r = client.put("/api/docs/big", headers=a, json={"title": "Big", "updated": 1, "data": {"blob": blob}})
    assert r.status_code == 413 and r.json()["detail"] == "This workbook is too large to sync (5 MB max)."
    bad = client.put("/api/docs/bad", headers=a, json={"title": "Bad"})
    assert bad.status_code == 422
