"""End-to-end API tests with the Anthropic HTTP layer mocked (no API key or network needed)."""

import io
import json

import anthropic
import httpx2
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routes import get_engine
from app.config import Settings
from app.main import app
from app.services.llm_engine import WorkbookEngine


def sse_response(payload: dict) -> httpx2.Response:
    text = json.dumps(payload)
    events = [
        ("message_start", {"type": "message_start", "message": {
            "id": "msg_1", "type": "message", "role": "assistant", "model": "claude-sonnet-5-5",
            "content": [], "stop_reason": None, "stop_sequence": None,
            "usage": {"input_tokens": 10, "output_tokens": 0}}}),
        ("content_block_start", {"type": "content_block_start", "index": 0,
                                 "content_block": {"type": "text", "text": ""}}),
        ("content_block_delta", {"type": "content_block_delta", "index": 0,
                                 "delta": {"type": "text_delta", "text": text}}),
        ("content_block_stop", {"type": "content_block_stop", "index": 0}),
        ("message_delta", {"type": "message_delta", "delta": {"stop_reason": "end_turn", "stop_sequence": None},
                           "usage": {"output_tokens": 50}}),
        ("message_stop", {"type": "message_stop"}),
    ]
    body = "".join(f"event: {e}\ndata: {json.dumps(d)}\n\n" for e, d in events)
    return httpx2.Response(200, headers={"content-type": "text/event-stream"}, content=body.encode())


def make_engine(replies: list[dict], seen: list[dict]) -> WorkbookEngine:
    queue = list(replies)

    def handler(request: httpx2.Request) -> httpx2.Response:
        seen.append({"headers": dict(request.headers), "body": json.loads(request.content)})
        return sse_response(queue.pop(0))

    client = anthropic.AsyncAnthropic(api_key="test", http_client=httpx2.AsyncClient(transport=httpx2.MockTransport(handler)))
    return WorkbookEngine(client=client, settings=Settings(repair_attempts=1))


def tiny_workbook(formula: str) -> dict:
    return {"title": "Tax", "sheets": [{"name": "Sales", "columns": [
        {"header": "Item"}, {"header": "Amount", "format": "currency"}, {"header": "Tax", "format": "currency"}],
        "rows": [{"cells": [{"value": "Widget"}, {"value": 100}, {"value": formula}]}]}]}


def test_generate_repairs_bad_reference_then_exports(budget_spec):
    seen: list[dict] = []
    bad = {"summary": "Built it.", "workbook": tiny_workbook("=Rates!B2*B2")}
    good = {"summary": "Built a sales sheet with 15% tax.", "workbook": tiny_workbook("=B2*0.15")}
    app.dependency_overrides[get_engine] = lambda: make_engine([bad, good], seen)
    try:
        client = TestClient(app)
        r = client.post("/api/workbooks/generate", json={"prompt": "Sales sheet with a 15% tax column"})
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["workbook"]["sheets"][0]["rows"][0]["cells"][2]["value"] == "=B2*0.15"
        assert data["summary"].startswith("Built a sales sheet")

        # The first answer was sent back with the error for a repair pass.
        assert len(seen) == 2
        repair_turn = seen[1]["body"]["messages"][-1]["content"]
        assert "Rates" in repair_turn and "does not exist" in repair_turn
        assert seen[0]["body"]["model"] == "claude-sonnet-5-5"
        assert seen[0]["body"]["fallbacks"] == "default"
        assert seen[0]["body"]["output_config"]["format"]["type"] == "json_schema"

        x = client.post("/api/workbooks/export", json=budget_spec.model_dump())
        assert x.status_code == 200
        assert 'filename="Small_Business_Budget_H1.xlsx"' in x.headers["content-disposition"]
        assert load_workbook(io.BytesIO(x.content)).sheetnames == ["Income", "Expenses", "Summary"]
    finally:
        app.dependency_overrides.clear()


def test_validate_endpoint_reports_and_fixes():
    client = TestClient(app)
    wb = tiny_workbook("=B2/A2")
    r = client.post("/api/workbooks/validate", json=wb)
    assert r.status_code == 200
    body = r.json()
    assert body["workbook"]["sheets"][0]["rows"][0]["cells"][2]["value"] == "=IFERROR(B2/A2,0)"
    assert body["issues"][0]["fixed"] is True
