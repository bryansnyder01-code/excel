import io

from openpyxl import load_workbook

from app.services.workbook_builder import workbook_to_bytes


def test_builds_valid_xlsx_with_formulas_and_styles(budget_spec):
    wb = load_workbook(io.BytesIO(workbook_to_bytes(budget_spec)))
    assert wb.sheetnames == ["Income", "Expenses", "Summary"]

    inc = wb["Income"]
    assert inc["A1"].value == "Source" and inc["A1"].font.bold
    assert inc["A1"].fill.fgColor.rgb.endswith("1F4E78")
    assert inc["H2"].value == "=SUM(B2:G2)"
    assert inc["B2"].number_format.startswith('"$"')
    assert inc.freeze_panes == "A2"
    assert inc["A5"].font.bold  # total row

    summ = wb["Summary"]
    assert summ["B2"].value == "='Income'!B5"
    assert summ["E2"].number_format == "0.0%"
    assert len(summ.conditional_formatting) == 2

    exp = wb["Expenses"]
    assert exp["J2"].value == '=IF(H2>10000,"High","Low")'
    assert exp.column_dimensions["A"].width >= len("Total expenses")


def test_manual_style_overrides(budget_spec):
    cell = budget_spec.sheets[0].rows[0].cells[0]
    cell.italic, cell.font_color, cell.fill_color = True, "#c00000", "fff2cc"
    cell.align, cell.border = "center", "thick"
    ws = load_workbook(io.BytesIO(workbook_to_bytes(budget_spec)))["Income"]
    a2 = ws["A2"]
    assert a2.font.italic and a2.font.color.rgb.endswith("C00000")
    assert a2.fill.fgColor.rgb.endswith("FFF2CC")
    assert a2.alignment.horizontal == "center"
    assert a2.border.left.style == "medium"
