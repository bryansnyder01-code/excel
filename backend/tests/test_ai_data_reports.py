"""=AI() cells, protection, dashboards, AI fill / eval, photo-to-table, live data and scheduled reports.

Claude is mocked at the HTTP layer (tests.test_api.make_engine); live data uses an httpx MockTransport
and a stubbed DNS resolver; email is captured, never sent.
"""

import asyncio
import base64
import io
import logging
import time
from datetime import datetime, timezone

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import data as data_api
from app.api import live
from app.api.accounts import get_store
from app.api.routes import get_engine
from app.config import get_settings
from app.main import app
from app.schemas import WorkbookSpec
from app.services import fetcher, reports as reports_service, scheduler
from app.services.fetcher import FetchError, FetchResult, gsheet_csv_url
from app.services.formula_validator import validate_workbook
from app.services.importer import import_file
from app.services.reports import get_report_store, key_numbers, next_send
from app.services.workbook_builder import workbook_to_bytes
from tests.test_api import make_engine
from tests.test_more import auth, signup

client = TestClient(app)

PNG = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64
PDF = b"%PDF-1.4\n%fake\n"


@pytest.fixture(autouse=True)
def _clean():
    yield
    app.dependency_overrides.clear()


@pytest.fixture
def no_recalc(monkeypatch):
    """Report tests don't start LibreOffice."""
    monkeypatch.setattr(reports_service, "recalculated_values", lambda spec, timeout=60: None)
    monkeypatch.setattr(scheduler, "recalculated_values", lambda spec: None)


def use_engine(replies, seen):
    app.dependency_overrides[get_engine] = lambda: make_engine(replies, seen)


def spec_with_extras() -> WorkbookSpec:
    return WorkbookSpec.model_validate({
        "title": "Sales",
        "sheets": [
            {"name": "Summary", "locked_ranges": ["A1:C1", "nonsense"],
             "columns": [{"header": "Region"}, {"header": "Revenue", "format": "currency", "currency": "EUR"},
                         {"header": "Mood"}],
             "rows": [
                 {"cells": [{"value": "East"}, {"value": 120}, {"value": '=AI("sentiment of", A2)', "ai_value": "Upbeat"}]},
                 {"cells": [{"value": "West"}, {"value": 80}, {"value": '=ai("sentiment", A3)'}]},
                 {"kind": "total", "cells": [{"value": "Total"}, {"value": "=SUM(B2:B3)"}]},
             ],
             "charts": [{"type": "column", "title": "Revenue", "categories": "A2:A3",
                         "series": [{"name": "Revenue", "range": "B2:B3"}]}]},
            {"name": "Inputs", "locked": True, "columns": [{"header": "Rate", "format": "percent"}],
             "rows": [{"cells": [{"value": 0.2}]}]},
        ],
        "dashboard": {"title": "Overview", "tiles": [
            {"id": "k1", "kind": "kpi", "title": "Total revenue", "ref": "Summary!B4", "compare_ref": "'Summary'!B2",
             "x": 0, "y": 0, "w": 3, "h": 1},
            {"id": "t1", "kind": "text", "title": "Notes", "text": "Q1 numbers", "x": 3, "y": 0, "w": 4, "h": 1},
            {"id": "c1", "kind": "chart", "title": "Revenue by region", "sheet": "Summary", "chart": 0,
             "x": 0, "y": 1, "w": 6, "h": 3},
        ]},
        "connections": [{"id": "x", "name": "Prices", "kind": "csv", "url": "https://example.com/p.csv", "sheet": "Prices"}],
    })


# ---------------------------------------------------------------------------
# Schema, builder, validator, importer
# ---------------------------------------------------------------------------


def test_old_workbooks_still_load():
    wb = WorkbookSpec.model_validate({"title": "Old", "sheets": [{"name": "A", "columns": [{"header": "x"}], "rows": []}]})
    assert wb.dashboard is None and wb.connections == [] and wb.sheets[0].locked is False
    assert wb.sheets[0].locked_ranges == []


def test_builder_writes_ai_values_protection_and_dashboard():
    wb = load_workbook(io.BytesIO(workbook_to_bytes(spec_with_extras())))
    assert wb.sheetnames == ["Dashboard", "Summary", "Inputs"]

    ws = wb["Summary"]
    assert ws["C2"].value == "Upbeat" and ws["C3"].value is None  # =AI() -> its last value, or empty
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws.protection.sheet
    assert ws["A1"].protection.locked and ws["C1"].protection.locked
    assert not ws["A2"].protection.locked and not ws["B4"].protection.locked
    assert wb["Inputs"].protection.sheet and wb["Inputs"]["A2"].protection.locked

    d = wb["Dashboard"]
    assert d["B1"].value == "Overview"
    assert d["B3"].value == "Total revenue"
    assert d["B4"].value == "='Summary'!B4"
    assert "€" in d["B4"].number_format
    assert "'Summary'!B2" in d["B5"].value  # comparison line
    assert d["E3"].value == "Notes" and d["E4"].value == "Q1 numbers"
    assert len(d._charts) == 1 and len(ws._charts) == 1


def test_validator_dashboard_and_locked_range_warnings():
    spec = spec_with_extras()
    spec.dashboard.tiles[0].ref = "Nope!B2"
    spec.dashboard.tiles[0].compare_ref = "Summary!Z99"
    spec.dashboard.tiles[2].chart = 3
    messages = [i.message for i in validate_workbook(spec)]
    assert any("Nope" in m and "does not exist" in m for m in messages)
    assert any("Summary!Z99" in m and "outside" in m for m in messages)
    assert any("chart 4" in m for m in messages)
    assert any('Locked range "nonsense"' in m for m in messages)
    assert not any("AI is not a recognised" in m for m in messages)


def test_validator_renames_follow_into_dashboard_and_connections():
    spec = spec_with_extras()
    spec.sheets[0].name = "Sum:mary"
    spec.dashboard.tiles[0].ref = "'Sum:mary'!B4"
    spec.dashboard.tiles[2].sheet = "Sum:mary"
    spec.connections[0].sheet = "Sum:mary"
    validate_workbook(spec)
    assert spec.sheets[0].name == "Sum-mary"
    assert spec.dashboard.tiles[0].ref == "'Sum-mary'!B4"
    assert spec.dashboard.tiles[2].sheet == "Sum-mary" and spec.connections[0].sheet == "Sum-mary"


def test_importer_reads_protection_back():
    spec = spec_with_extras()
    spec.dashboard = None
    spec.sheets[0].locked_ranges = ["A1:C1", "B3:B4"]
    wb, _ = import_file("s.xlsx", workbook_to_bytes(spec))
    summary, inputs = wb.sheets
    assert summary.locked is False and summary.locked_ranges == ["A1:C1", "B3:B4"]
    assert inputs.locked is True and inputs.locked_ranges == []


# ---------------------------------------------------------------------------
# AI fill and =AI() evaluation
# ---------------------------------------------------------------------------


def test_ai_fill_numbers_in_order_and_counts_one_request():
    seen: list[dict] = []
    reply = {"values": [{"row": 1, "value": "$1,200"}, {"row": 0, "value": "15%"}, {"row": 2, "value": "n/a"}]}
    use_engine([reply], seen)
    r = client.post("/api/workbooks/ai-fill", json={
        "instruction": "Estimated price", "headers": ["Item"], "rows": [["Desk"], ["Laptop"], ["Air"]], "output": "number"})
    assert r.status_code == 200, r.text
    assert r.json() == {"values": [0.15, 1200.0, None]}
    body = seen[0]["body"]
    assert body["output_config"]["effort"] == "low" and body["max_tokens"] <= 16000
    assert "Laptop" in body["messages"][0]["content"]
    assert client.get("/api/workbooks/usage").json()["used"] == 1


def test_ai_fill_text_and_limits():
    seen: list[dict] = []
    use_engine([{"values": [{"row": 0, "value": " Furniture "}]}], seen)
    r = client.post("/api/workbooks/ai-fill", json={"instruction": "Category", "headers": ["Item"],
                                                    "rows": [["Desk"], ["Pen"]], "output": "text"})
    assert r.json() == {"values": ["Furniture", None]}  # a row the model skipped comes back empty

    r = client.post("/api/workbooks/ai-fill", json={"instruction": "x", "headers": [], "rows": [["a"]] * 201})
    assert r.status_code == 422 and "200 rows" in r.json()["detail"]
    assert len(seen) == 1


def test_ai_eval_batch():
    seen: list[dict] = []
    use_engine([{"results": [{"key": "Sheet1!C2", "value": "Positive"}, {"key": "bogus", "value": "x"}]}], seen)
    r = client.post("/api/workbooks/ai-eval", json={"items": [
        {"key": "Sheet1!C2", "prompt": "Sentiment?", "inputs": ["Love it"]},
        {"key": "Sheet1!C3", "prompt": "Sentiment?", "inputs": ["Meh", 3]},
        {"key": "Sheet1!C4", "prompt": "  ", "inputs": []},
    ]})
    assert r.status_code == 200, r.text
    assert r.json() == {"results": {"Sheet1!C2": "Positive", "Sheet1!C3": "", "Sheet1!C4": ""}}
    text = seen[0]["body"]["messages"][0]["content"]
    assert "Love it" in text and "Sheet1!C4" not in text  # empty prompts aren't sent

    r = client.post("/api/workbooks/ai-eval",
                    json={"items": [{"key": str(i), "prompt": "x", "inputs": []} for i in range(101)]})
    assert r.status_code == 422


# ---------------------------------------------------------------------------
# Photo / PDF to table
# ---------------------------------------------------------------------------


EXTRACTED = {
    "summary": "Found one receipt.", "title": "Receipt",
    "tables": [{"name": "Lines", "columns": [
        {"header": "Date", "format": "date"}, {"header": "Item", "format": "text"},
        {"header": "Price", "format": "currency", "currency": "eur"}],
        "rows": [{"cells": ["2026/03/05", "Coffee", "€3.50"]},
                 {"cells": ["2026-03-05", "=HYPERLINK(\"x\")", 12]},
                 {"kind": "total", "cells": [None, "Total", 15.5]}]}],
}


def test_extract_image_to_workbook():
    seen: list[dict] = []
    use_engine([EXTRACTED], seen)
    r = client.post("/api/workbooks/extract", files={"file": ("receipt.png", PNG, "image/png")},
                    data={"hint": "a cafe receipt"})
    assert r.status_code == 200, r.text
    out = r.json()
    sheet = out["workbook"]["sheets"][0]
    assert sheet["name"] == "Lines" and sheet["columns"][2]["currency"] == "EUR"
    rows = [[c["value"] for c in row["cells"]] for row in sheet["rows"]]
    assert rows[0] == ["2026-03-05", "Coffee", 3.5]
    assert rows[1][1].startswith("'=")  # never a formula
    assert sheet["rows"][2]["kind"] == "total"
    assert out["summary"].startswith("Found one receipt.")

    content = seen[0]["body"]["messages"][0]["content"]
    assert content[0]["type"] == "image" and content[0]["source"]["media_type"] == "image/png"
    assert base64.b64decode(content[0]["source"]["data"]) == PNG
    assert "a cafe receipt" in content[1]["text"]
    assert client.get("/api/workbooks/usage").json()["used"] == 1


def test_extract_pdf_uses_document_block_and_rejects_bad_files():
    seen: list[dict] = []
    use_engine([EXTRACTED], seen)
    r = client.post("/api/workbooks/extract", files={"file": ("scan.pdf", PDF, "application/pdf")})
    assert r.status_code == 200, r.text
    assert seen[0]["body"]["messages"][0]["content"][0]["type"] == "document"

    r = client.post("/api/workbooks/extract", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert r.status_code == 415 and "PDF" in r.json()["detail"]
    r = client.post("/api/workbooks/extract",
                    files={"file": ("big.pdf", PDF + b"0" * (10 * 1024 * 1024), "application/pdf")})
    assert r.status_code == 413
    assert len(seen) == 1


def test_extract_with_no_tables_is_422():
    use_engine([{"summary": "Nothing.", "title": "x", "tables": []}], [])
    r = client.post("/api/workbooks/extract", files={"file": ("a.png", PNG, "image/png")})
    assert r.status_code == 422 and "No tables" in r.json()["detail"]


# ---------------------------------------------------------------------------
# Live data
# ---------------------------------------------------------------------------


@pytest.fixture
def public_dns(monkeypatch):
    hosts = {"internal.example": "10.0.0.5"}

    async def resolve(host, port):
        return [hosts.get(host, "93.184.216.34")]

    monkeypatch.setattr(fetcher, "resolve_host", resolve)
    monkeypatch.setattr(data_api, "fetch_limiter", data_api.RateLimiter(60, 3600))
    return hosts


def mock_http(routes: dict, seen: list | None = None):
    def handler(request: httpx.Request) -> httpx.Response:
        if seen is not None:
            seen.append(str(request.url))
        status, headers, body = routes.get(str(request.url), (404, {}, b"missing"))
        return httpx.Response(status, headers=headers, content=body)

    async def dep():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as c:
            yield c

    app.dependency_overrides[data_api.fetch_http] = dep


def test_fetch_csv(public_dns):
    mock_http({"https://example.com/p.csv": (200, {"content-type": "text/csv"},
                                             b"Item,Price,Date\nDesk,\"$1,200.50\",2026-01-02\nPen,3,\n")})
    r = client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/p.csv"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["columns"] == ["Item", "Price", "Date"]
    assert body["rows"] == [["Desk", 1200.5, "2026-01-02"], ["Pen", 3.0, None]]
    assert abs(body["fetched_at"] - time.time() * 1000) < 60_000


def test_fetch_gsheet_link_conversion_and_private_sheet(public_dns):
    url = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/edit#gid=42"
    export = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz012345/export?format=csv&gid=42"
    assert gsheet_csv_url(url) == export
    assert gsheet_csv_url(url.replace("#gid=42", "?usp=sharing")).endswith("gid=0")
    with pytest.raises(FetchError):
        gsheet_csv_url("https://example.com/spreadsheets/d/abc")

    seen: list[str] = []
    mock_http({export: (307, {"location": "https://accounts.google.com/login"}, b""),
               "https://accounts.google.com/login": (200, {"content-type": "text/html"}, b"<html>sign in</html>")}, seen)
    r = client.post("/api/data/fetch", json={"kind": "gsheet", "url": url})
    assert r.status_code == 422 and "Anyone with the link" in r.json()["detail"]
    assert seen == [export, "https://accounts.google.com/login"]


def test_fetch_json_path_union_of_keys(public_dns):
    mock_http({"https://api.example.com/x": (200, {"content-type": "application/json"},
                                             b'{"data": {"items": [{"a": 1, "b": "x"}, {"c": {"d": 1}, "a": true}]}}')})
    r = client.post("/api/data/fetch", json={"kind": "json", "url": "https://api.example.com/x", "path": "data.items"})
    assert r.status_code == 200, r.text
    assert r.json()["columns"] == ["a", "b", "c"]
    assert r.json()["rows"] == [[1.0, "x", None], ["TRUE", None, '{"d": 1}']]

    r = client.post("/api/data/fetch", json={"kind": "json", "url": "https://api.example.com/x"})
    assert r.status_code == 422 and "data.items" in r.json()["detail"]


@pytest.mark.parametrize("url", [
    "http://127.0.0.1/x", "http://10.1.2.3/", "http://169.254.169.254/latest/meta-data", "http://[::1]/",
    "http://internal.example/data.csv", "ftp://example.com/x.csv", "file:///etc/passwd", "http://0.0.0.0/",
])
def test_fetch_refuses_private_and_odd_addresses(public_dns, url):
    seen: list[str] = []
    mock_http({}, seen)
    r = client.post("/api/data/fetch", json={"kind": "csv", "url": url})
    assert r.status_code == 422
    assert seen == []  # never requested


def test_fetch_redirect_checks(public_dns):
    seen: list[str] = []
    mock_http({"https://example.com/a": (302, {"location": "http://internal.example/secret"}, b"")}, seen)
    r = client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/a"})
    assert r.status_code == 422 and "private" in r.json()["detail"]
    assert seen == ["https://example.com/a"]

    loop = {f"https://example.com/{i}": (302, {"location": f"/{i + 1}"}, b"") for i in range(10)}
    mock_http(loop)
    r = client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/0"})
    assert r.status_code == 422 and "redirects too many" in r.json()["detail"]


def test_fetch_size_cap_and_rate_limit(public_dns, monkeypatch):
    mock_http({"https://example.com/big.csv": (200, {}, b"a\n" + b"1\n" * (3 * 1024 * 1024))})
    r = client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/big.csv"})
    assert r.status_code == 422 and "5 MB" in r.json()["detail"]

    monkeypatch.setattr(data_api, "fetch_limiter", data_api.RateLimiter(1, 3600))
    mock_http({"https://example.com/p.csv": (200, {}, b"a\n1\n")})
    assert client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/p.csv"}).status_code == 200
    assert client.post("/api/data/fetch", json={"kind": "csv", "url": "https://example.com/p.csv"}).status_code == 429


# ---------------------------------------------------------------------------
# Scheduled refresh
# ---------------------------------------------------------------------------


def put_doc(token: str, doc_id: str, workbook: dict) -> None:
    r = client.put(f"/api/docs/{doc_id}", headers=auth(token),
                   json={"title": workbook["title"], "updated": 1, "data": {"workbook": workbook, "title": workbook["title"]}})
    assert r.status_code == 200, r.text


def connected_workbook(last=None) -> dict:
    return {"title": "Prices", "sheets": [{"name": "Main", "columns": [{"header": "x"}], "rows": []},
                                          {"name": "Feed", "columns": [{"header": "Price", "width": 22}], "rows": []}],
            "connections": [{"id": "c1", "name": "Feed", "kind": "csv", "url": "https://example.com/p.csv",
                             "sheet": "Feed", "refresh_minutes": 15, "last_refreshed": last},
                            {"id": "c2", "name": "Manual", "kind": "json", "url": "https://example.com/j",
                             "sheet": "New tab", "refresh_minutes": None}]}


def fake_fetch(calls: list):
    async def fetch(kind, url, path=None):
        calls.append((kind, url, path))
        return FetchResult(["Item", "Price"], [["Desk", 10.0], ["Pen", None]],
                           [fetcher.NumberFormat.text, fetcher.NumberFormat.currency], fetched_at=1)
    return fetch


def test_scheduled_refresh_saves_due_connections():
    alice = signup("refresh@example.com")
    put_doc(alice["token"], "doc-r", connected_workbook())
    put_doc(alice["token"], "doc-fresh", connected_workbook(last=int(time.time() * 1000)))
    source = scheduler.StoreDocSource(get_store(), get_settings().data_dir)
    scheduler._attempts.clear()
    calls: list = []
    n = asyncio.run(scheduler.refresh_due_connections(source, fake_fetch(calls)))
    assert n == 1 and calls == [("csv", "https://example.com/p.csv", None)]

    doc = client.get("/api/docs/doc-r", headers=auth(alice["token"])).json()
    wb = doc["data"]["workbook"]
    feed = wb["sheets"][1]
    assert [c["header"] for c in feed["columns"]] == ["Item", "Price"]
    assert feed["columns"][1]["format"] == "currency"
    assert feed["rows"][0]["cells"] == [{"value": "Desk"}, {"value": 10.0}]
    assert wb["connections"][0]["last_refreshed"] > 0 and wb["connections"][1].get("last_refreshed") is None

    # A failed source backs off instead of retrying every minute.
    async def broken(kind, url, path=None):
        calls.append("broken")
        raise FetchError("down")
    put_doc(alice["token"], "doc-b", connected_workbook())
    asyncio.run(scheduler.refresh_doc(source, "doc-b", broken))
    asyncio.run(scheduler.refresh_doc(source, "doc-b", broken))
    assert calls.count("broken") == 1


def test_scheduled_refresh_goes_through_open_live_room():
    alice = signup("live-refresh@example.com")
    put_doc(alice["token"], "doc-live", connected_workbook())

    class FakeRoom:
        def __init__(self, workbook):
            self.workbook, self.version, self.sent = workbook, 3, []

        def changed(self):
            self.version += 1

        def broadcast(self, message, exclude=None):
            self.sent.append(message)

    room = FakeRoom(connected_workbook() | {"title": "Edited live"})
    live.rooms["doc-live"] = room
    try:
        source = scheduler.StoreDocSource(get_store(), get_settings().data_dir)
        scheduler._attempts.clear()
        n = asyncio.run(scheduler.refresh_doc(source, "doc-live", fake_fetch([]), force=True))
    finally:
        live.rooms.pop("doc-live", None)
    assert n == 2  # forced: the manual JSON connection too, into a new sheet
    assert room.version == 4 and room.workbook["title"] == "Edited live"
    (msg,) = room.sent
    assert msg["type"] == "snapshot" and msg["from"] == "system" and msg["version"] == 4
    assert [s["name"] for s in msg["workbook"]["sheets"]] == ["Main", "Feed", "New tab"]
    assert msg["workbook"]["sheets"][1]["columns"][0] == {"header": "Item", "format": "text"}


# ---------------------------------------------------------------------------
# Scheduled email reports
# ---------------------------------------------------------------------------


def report_workbook() -> dict:
    return {"title": "Budget", "sheets": [{"name": "Costs",
            "columns": [{"header": "Item"}, {"header": "Cost", "format": "currency"}],
            "rows": [{"cells": [{"value": "Rent"}, {"value": 1200}]},
                     {"cells": [{"value": "Power"}, {"value": 300}]},
                     {"kind": "total", "cells": [{"value": "Total"}, {"value": "=SUM(B2:B3)"}]}]}]}


def test_next_send_schedules():
    t = datetime(2026, 10, 19, 10, 30, tzinfo=timezone.utc).timestamp()  # a Monday
    at = lambda s: datetime.fromtimestamp(s, timezone.utc).strftime("%Y-%m-%d %H:%M")  # noqa: E731
    assert at(next_send("daily", None, None, 9, t)) == "2026-10-20 09:00"
    assert at(next_send("daily", None, None, 11, t)) == "2026-10-19 11:00"
    assert at(next_send("weekly", 0, None, 9, t)) == "2026-10-26 09:00"
    assert at(next_send("weekly", 4, None, 9, t)) == "2026-10-23 09:00"
    assert at(next_send("monthly", None, 19, 11, t)) == "2026-10-19 11:00"
    assert at(next_send("monthly", None, 1, 8, t)) == "2026-11-01 08:00"
    dec = datetime(2026, 12, 20, tzinfo=timezone.utc).timestamp()
    assert at(next_send("monthly", None, 5, 0, dec)) == "2027-01-05 00:00"


def test_reports_crud_and_access(no_recalc):
    alice, bob = signup("rep-owner@example.com"), signup("rep-other@example.com")
    put_doc(alice["token"], "doc-rep", report_workbook())
    base = "/api/docs/doc-rep/reports"
    good = {"frequency": "weekly", "weekday": 2, "hour": 7, "recipients": ["A@Example.com", "a@example.com"],
            "sheet": "Costs", "attach": "xlsx", "include_summary": False}
    r = client.post(base, headers=auth(alice["token"]), json=good)
    assert r.status_code == 200, r.text
    rep = r.json()
    assert rep["recipients"] == ["a@example.com"] and rep["weekday"] == 2 and rep["day"] is None
    assert rep["last_sent"] is None and rep["next_send"] > time.time() * 1000
    assert datetime.fromtimestamp(rep["next_send"] / 1000, timezone.utc).weekday() == 2

    for bad in ({"recipients": ["nope"]}, {"recipients": []}, {"recipients": [f"p{i}@x.com" for i in range(11)]},
                {"sheet": "Missing"}, {"hour": 24}, {"frequency": "hourly"}):
        assert client.post(base, headers=auth(alice["token"]), json=good | bad).status_code == 422, bad

    assert client.get(base, headers=auth(bob["token"])).status_code == 404
    assert client.post(base, headers=auth(bob["token"]), json=good).status_code == 404
    assert client.get(base).status_code == 401

    # An editor (joined through the invite link) can manage reports too.
    token = client.post("/api/docs/doc-rep/invite", headers=auth(alice["token"])).json()["token"]
    assert client.post(f"/api/invites/{token}/accept", headers=auth(bob["token"])).status_code == 200
    assert [x["id"] for x in client.get(base, headers=auth(bob["token"])).json()] == [rep["id"]]

    assert client.delete(f"{base}/{rep['id']}", headers=auth(bob["token"])).status_code == 204
    assert client.get(base, headers=auth(alice["token"])).json() == []
    assert client.delete(f"{base}/{rep['id']}", headers=auth(alice["token"])).status_code == 404


def test_send_now_logs_without_smtp_and_uses_owner_ai_allowance(no_recalc, caplog):
    alice = signup("rep-send@example.com")
    put_doc(alice["token"], "doc-send", report_workbook())
    seen: list[dict] = []
    use_engine([{"summary": "Rent dominates costs."}], seen)
    rep = client.post("/api/docs/doc-send/reports", headers=auth(alice["token"]), json={
        "frequency": "daily", "hour": 6, "recipients": ["boss@example.com"], "attach": "none",
        "include_summary": True}).json()
    with caplog.at_level(logging.WARNING, logger="app.services.scheduler"):
        r = client.post(f"/api/docs/doc-send/reports/{rep['id']}/send", headers=auth(alice["token"]))
    assert r.status_code == 204, r.text
    log_text = caplog.text
    assert "Budget — daily report" in log_text and "boss@example.com" in log_text
    assert "Rent dominates costs." in log_text
    assert "Total: Cost: =SUM(B2:B3)" in log_text  # no formula engine here: shown honestly as a formula
    assert "Open the workbook to see the calculated totals" in log_text
    assert "/?doc=doc-send" in log_text
    assert "=SUM(B2:B3)" in seen[0]["body"]["messages"][0]["content"]
    usage = client.get("/api/workbooks/usage", headers=auth(alice["token"])).json()
    assert usage["used"] == 1
    after = client.get("/api/docs/doc-send/reports", headers=auth(alice["token"])).json()[0]
    assert after["last_sent"] is not None and after["next_send"] == rep["next_send"]


def test_due_reports_are_emailed_with_attachment_and_rescheduled(no_recalc, monkeypatch):
    alice = signup("rep-due@example.com")
    put_doc(alice["token"], "doc-due", report_workbook())
    rep = client.post("/api/docs/doc-due/reports", headers=auth(alice["token"]), json={
        "frequency": "monthly", "day": 3, "hour": 9, "recipients": ["x@example.com", "y@example.com"],
        "sheet": "Costs", "attach": "xlsx", "include_summary": False}).json()

    s = get_settings()
    monkeypatch.setattr(s, "smtp_host", "smtp.example.com")
    monkeypatch.setattr(s, "smtp_from", "reports@example.com")
    sent: list = []
    monkeypatch.setattr(scheduler, "send_mail", lambda settings, to, subject, body, attachments=None:
                        sent.append((to, subject, body, attachments)))
    values = {"Costs": [["Item", "Cost"], ["Rent", 1200], ["Power", 300], ["Total", 1500]]}
    monkeypatch.setattr(scheduler, "recalculated_values", lambda spec: values)

    source = scheduler.StoreDocSource(get_store(), s.data_dir)
    later = rep["next_send"] / 1000 + 5
    n = asyncio.run(scheduler.send_due_reports(source, get_report_store(), None, s, None, now=later))
    assert n == 1 and [x[0] for x in sent] == ["x@example.com", "y@example.com"]
    to, subject, body, attachments = sent[0]
    assert subject == "Budget — monthly report"
    assert "1,500" in body and "=SUM" not in body  # calculated values from the recalculated grid
    ((name, data, mime),) = attachments
    assert name == "Budget.xlsx" and mime.endswith("spreadsheetml.sheet")
    assert load_workbook(io.BytesIO(data))["Costs"]["B4"].value == "=SUM(B2:B3)"

    after = client.get("/api/docs/doc-due/reports", headers=auth(alice["token"])).json()[0]
    assert after["last_sent"] == int(later * 1000) and after["next_send"] > rep["next_send"]
    # Not due again until next month.
    assert asyncio.run(scheduler.send_due_reports(source, get_report_store(), None, s, None, now=later + 60)) == 0


def test_key_numbers_without_totals_shows_first_rows():
    spec = WorkbookSpec.model_validate({"title": "T", "sheets": [{"name": "Log", "columns": [{"header": "Day"}, {"header": "Mood"}],
        "rows": [{"cells": [{"value": f"D{i}"}, {"value": '=AI("x")', "ai_value": "ok"}]} for i in range(15)]}]})
    text, uncalculated = key_numbers(spec, None, None)
    assert "D0" in text and "D9" in text and "D10" not in text and "first 10 of 15 rows" in text
    assert "ok" in text and not uncalculated


def test_scheduled_refresh_publishes_through_a_brokered_room():
    alice = signup("broker-refresh@example.com")
    put_doc(alice["token"], "doc-brk", connected_workbook())

    class BrokeredRoom:
        broker = object()

        def __init__(self, workbook):
            self.workbook, self.events = workbook, []

        async def publish(self, event):
            self.events.append(event)

    room = BrokeredRoom(connected_workbook())
    live.rooms["doc-brk"] = room
    try:
        source = scheduler.StoreDocSource(get_store(), get_settings().data_dir)
        scheduler._attempts.clear()
        assert asyncio.run(scheduler.refresh_doc(source, "doc-brk", fake_fetch([]))) == 1
    finally:
        live.rooms.pop("doc-brk", None)
    (event,) = room.events
    assert event["kind"] == "snapshot" and event["from"] == "system" and event["conn"] is None
    assert event["workbook"]["sheets"][1]["rows"][0]["cells"][0] == {"value": "Desk"}


def test_a_due_report_is_claimed_once():
    store = get_report_store()
    rep = store.create("doc-x", "u1", {"frequency": "daily", "weekday": None, "day": None, "hour": 3,
                                       "recipients": ["a@b.co"], "sheet": None, "attach": "none",
                                       "include_summary": False})
    now = rep["next_send"] / 1000 + 1
    assert store.claim(rep["id"], rep["next_send"], now) is True
    assert store.claim(rep["id"], rep["next_send"], now) is False


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice not installed")
def test_recalculated_values_with_libreoffice():
    spec = WorkbookSpec.model_validate(report_workbook())
    values = reports_service.recalculated_values(spec)
    assert values["Costs"][3] == ["Total", 1500]
    text, uncalculated = key_numbers(spec, None, values)
    assert "Total: Cost: 1,500" in text and not uncalculated


def test_lifespan_starts_and_stops_background_jobs(monkeypatch):
    started = []

    async def fake_loop(engine_factory, interval=60):
        started.append(engine_factory)
        await asyncio.sleep(3600)

    monkeypatch.setattr(scheduler, "run_forever", fake_loop)
    monkeypatch.setattr(get_settings(), "background_jobs", True)
    with TestClient(app) as c:
        assert c.get("/health").status_code == 200
    assert len(started) == 1
    monkeypatch.setattr(get_settings(), "background_jobs", False)
    with TestClient(app):
        pass
    assert len(started) == 1
